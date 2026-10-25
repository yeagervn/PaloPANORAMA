#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
run_all_features.py - CHAY THU TOAN BO TINH NANG ban PANORAMA (ca 3 tab), KHONG
CAN GUI, doi chieu voi ky vong va in ra bang PASS/FAIL.

Day la "ban nghiem thu tu dong": no goi dung cac ham ma GUI goi (cung pipeline,
cung exporter), chi khac la khong ve giao dien - nen ket qua PASS o day nghia la
phan LOI (nghiep vu + API) da chay dung tu dau den cuoi.

CACH CHAY (2 buoc):
    # cua so 1
    python testlab/mock_panorama_api.py --port 8443
    # cua so 2
    python testlab/run_all_features.py --host 127.0.0.1:8443

Hoac de script tu bat/tat mock (1 lenh duy nhat):
    python testlab/run_all_features.py --auto

File ket qua xuat ra: testlab/output/
Exit code: 0 = tat ca PASS, 1 = co FAIL.
"""

import argparse
import contextlib
import datetime
import os
import shutil
import subprocess
import sys
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
OUT = os.path.join(HERE, "output")

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((bool(cond), name, detail))
    mark = "PASS" if cond else "FAIL"
    print(f"   [{mark}] {name}" + (f"   -- {detail}" if detail else ""))
    return bool(cond)


def section(title):
    print()
    print("-" * 78)
    print(f"  {title}")
    print("-" * 78)


# ======================================================================= TAB 1
DG = "DG-HN"
SERIAL = "001122334455"


def tab_check_rule(host, key):
    section(f"TAB 1 - CHECK RULE  (Device Group={DG}, Target={SERIAL})")
    from check_rule_pan.core import (fetch_panorama_config_xml, load_pan_config_from_xml_string,
                                      fetch_candidate_config_xml, parse_candidate_config,
                                      diff_candidate_config, get_panorama_devicegroups,
                                      get_panorama_devices, devicegroup_devices,
                                      list_devicegroups)
    import xml.etree.ElementTree as ET
    from check_rule_pan.resolver import GroupResolver, ServiceResolver
    from check_rule_pan.index import build_device_index, build_zone_lookup
    from check_rule_pan.matcher import build_rule_infos, classify_row, apply_candidate_overlay
    from check_rule_pan.request import parse_request_file
    from check_rule_pan.exporter import (export_results, _compute_action_level,
                                          build_name_rule_cell)

    xml = fetch_panorama_config_xml(host, key)
    root_el = ET.fromstring(xml)
    check("0.1 - API doc config Panorama", len(xml) > 5000, f"{len(xml)} ky tu")
    check("0.1 - op-command liet ke Device Group",
          get_panorama_devicegroups(host, key) == ["DG-ALL", "DG-DC", "DG-HN"],
          str(get_panorama_devicegroups(host, key)))
    check("0.1 - op-command liet ke firewall (de chon Target)",
          [s for s, _h in get_panorama_devices(host, key)] == ["001122334455", "009988776655"],
          str(get_panorama_devices(host, key)))
    check("0.1 - doc DG thang tu config (duong du phong)",
          list_devicegroups(root_el) == ["DG-ALL", "DG-DC", "DG-HN"])

    cfg = load_pan_config_from_xml_string(xml, DG, SERIAL)
    # Assert nay truoc day viet qua sat chu (`[:2] == [...]`) nen chi can them 1
    # rule o Shared-PRE la vo, du thu tu lap rap khong he sai. Nay khang dinh
    # CHUOI vi_tri PHAN BIET - dung y dinh that cua 0.3, va CHAT HON (kiem ca 8
    # buoc, khong chi 2 dau + 1 cuoi).
    _thu_tu = list(dict.fromkeys(r.vi_tri for r in cfg["rules"]))
    check("0.3 - lap rap rulebase theo thu tu PAN-OS danh gia",
          _thu_tu == ["Shared-PRE", "DG-ALL-PRE", "DG-DC-PRE", "DG-HN-PRE",
                      "DG-HN-POST", "DG-DC-POST", "DG-ALL-POST", "Shared-POST"],
          " -> ".join(_thu_tu))
    check("0.3 - POST-rulebase NGUOC chieu PRE",
          [r.dg for r in cfg["rules"] if r.rulebase == "post" and r.dg != "shared"]
          == ["DG-HN", "DG-DC", "DG-ALL"])
    check("0.2 - chuoi thua ke Device Group",
          cfg["devicegroup_chain"] == ["shared", "DG-ALL", "DG-DC", "DG-HN"],
          " -> ".join(cfg["devicegroup_chain"]))
    check("0.2 - object DG gan hon DE object DG xa hon/Shared",
          cfg["address_objects"]["OVERRIDE-ME"].display == "10.3.3.3/32",
          cfg["address_objects"]["OVERRIDE-ME"].display)
    check("0.2 - thua ke object tu DG cha + Shared",
          cfg["address_objects"]["RANGE-WORKER"].display == "10.10.10.100-10.10.10.110"
          and "WEB-FARM" in cfg["address_objects"])
    check("0.6 - Zone lay tu Template cua Target Firewall",
          cfg["template_used"] == ["TPL-BASE"] and len(cfg["zones"]) == 3,
          f"template={cfg['template_used']} zones={list(cfg['zones'])}")
    check("0.1 - kiem tra cheo Target co thuoc DG dang chon khong",
          SERIAL in devicegroup_devices(root_el, DG),
          str(devicegroup_devices(root_el, DG)))

    cfg_nodg = load_pan_config_from_xml_string(xml, "DG-ALL", SERIAL)
    check("0.2/0.3 - doi DG thi pham vi doi theo",
          len(cfg_nodg["rules"]) < len(cfg["rules"]),
          f"DG-ALL: {len(cfg_nodg['rules'])} rule / {DG}: {len(cfg['rules'])} rule")

    cfg_notarget = load_pan_config_from_xml_string(xml, DG)
    check("0.6 - khong co Target -> Zone TU VO HIEU im lang",
          cfg_notarget["zones"] == {} and cfg_notarget["interface_subnets"] == [])

    gr = GroupResolver(cfg["address_objects"], cfg["address_groups"])
    sr = ServiceResolver(cfg["service_objects"], cfg["service_groups"])
    idx = build_device_index(cfg, gr, sr)
    ris = build_rule_infos(cfg["rules"], gr, sr)
    # B1: lab phai chay DUNG nhu GUI - tuc co truyen static route tu Template.
    zl = build_zone_lookup(cfg["interface_subnets"], cfg["zones"],
                            routes=cfg.get("static_routes") or None)
    # B1: bang nay nay gom CA dai suy tu static route trong Template (3 interface
    # + 2 route) nen con so dung la 5. Ky vong cu "== 3" la truoc khi co B1.
    check("0.6 - dung duoc bang tra cuu Zone tu Template (interface + static route)",
          len(zl) == 5, f"{len(zl)} dai")

    cand_ris, cand_idx = [], None
    try:
        cand_cfg = parse_candidate_config(fetch_candidate_config_xml(host, key), DG)
        diff = diff_candidate_config(cfg, cand_cfg)
        cgr = GroupResolver(cand_cfg["address_objects"], cand_cfg["address_groups"])
        csr = ServiceResolver(cand_cfg["service_objects"], cand_cfg["service_groups"])
        cand_idx = build_device_index(diff, cgr, csr)
        cand_ris = build_rule_infos(diff["rules"], cgr, csr)
        check("0.4 - candidate-config doc theo cau truc DG + loc ra phan THAT SU moi",
              len(diff["rules"]) == 1 and "HOST-NEW-APP" in diff["address_objects"],
              f"{len(diff['rules'])} rule moi, {len(diff['address_objects'])} object moi")
    except Exception as e:
        check("0.4 - candidate-config", False, str(e))

    req = parse_request_file(os.path.join(HERE, "samples", "request_check_rule.xlsx"))
    results = []
    for row in req.rows:
        r = classify_row(row, ris, idx, zone_lookup=zl)
        if r.tier != "LIVE_MATCH":
            apply_candidate_overlay(r, row, cand_ris, cand_idx)
        results.append(r)

    def row_of(i):
        return results[i], _compute_action_level(results[i])

    EXPECT = [
        (0, "TOP1", "none", "FULL MATCH"),
        (1, "TOP1.1", "add", "chi thieu Port"),
        (2, "TOP2", "add", "thieu Source"),
        (3, "TOP3", "add", "thieu Dest"),
        (4, "TOP1", "none", "range khop ip-range object"),
        (5, "TOP1", "none", "range RUT GON octet cuoi"),
        (6, "TOP1", "none", "range dung EN-DASH"),
        (7, "TOP1", "none", "range = HOP 2 object"),
        (8, None, "wide_full_match", "cover qua subnet rong"),
        (10, "TOP1", "add", "khop rule nhung bi DENY chan truoc -> shadow"),
        (11, "SAME_SUBNET", "same_subnet", "cung subnet"),
        (13, "TOP4", None, "chua co gi"),
        (14, None, "add", "token rac"),
    ]
    for i, tier, act, label in EXPECT:
        r, a = row_of(i)
        okc = (tier is None or r.tier == tier) and (act is None or a == act)
        check(f"Dong {i + 1}: {label}", okc, f"tier={r.tier} action={a}")

    check("Dong 11: SHADOW boi rule deny o Shared-PRE (dung thu tu lap rap 0.3)",
          results[10].shadow_warning == "S00-BLOCK-BADNET", str(results[10].shadow_warning))
    # Dong 10: KY VONG CU ("phai co zone_warning") da LAC HAU sau cong chan zone
    # 10/2026 - va cai moi DUNG HON: D03-ZONE-GAP bi loai vi lech zone, tool tim ra
    # D03B-ZONE-ANY-OK (zone any) la rule THUC SU cho phep luong nay. Doi assert
    # (khong xoa) y nhu da lam o ban standalone.
    _r10 = results[9].rule or results[9].nearest_rule
    check("Dong 10: KHONG chon rule lech zone (D03-ZONE-GAP)",
          _r10 is not None and _r10.name != "D03-ZONE-GAP", _r10 and _r10.name)
    check("Dong 10: chon dung rule zone-any lam rule cho phep that",
          _r10 is not None and _r10.name == "D03B-ZONE-ANY-OK"
          and results[9].zone_verdict == "match",
          f"{_r10 and _r10.name} verdict={results[9].zone_verdict}")
    # --- P2 (10/2026, CHI CO O BAN PANORAMA): deny o Shared-PRE ---
    # Thu tu lap rap 0.3: Shared-PRE -> DG-chain-PRE -> DG-chain-POST -> Shared-POST.
    # KHONG rule nao trong Device Group co the dung TRUOC Shared-PRE, nghia la "them
    # rule moi vao DG" o truong hop nay la VO NGHIA o moi vi tri. Ket luan dung phai
    # la "phai sua Shared-PRE" - do la thong tin ma khong phien ban standalone nao
    # can, va la gia tri rieng cua ban Panorama.
    _r16, _a16 = row_of(15)
    check("Dong 16: P2 - nhan dien deny nam o SHARED-PRE",
          _r16.blocking_deny_rule == "S00-BLOCK-BADNET" and _r16.blocked_at_shared_pre,
          f"deny={_r16.blocking_deny_rule} shared_pre={_r16.blocked_at_shared_pre}")
    check("Dong 16: P2 - ket luan la KHONG THE cap phep tu Device Group nay",
          "KHONG THE cap phep" in (_r16.blocking_deny_detail or "")
          and "Shared-PRE" in (_r16.blocking_deny_detail or ""),
          str(_r16.blocking_deny_detail)[:90])
    _c16 = "".join(getattr(b, "text", "") or ""
                   for b in build_name_rule_cell(_r16)[0])
    check("Dong 16: P2 - o Name Rule noi dung ban chat, khong phai 'dat len tren'",
          "KHONG THE CAP PHEP TU DEVICE GROUP NAY" in _c16.upper()
          and "phai dat TREN" not in _c16,
          "nhan rieng cho Shared-PRE")

    # --- B1 (10/2026): zone suy tu BANG DINH TUYEN trong Template ---
    # Dong 17: nguon 203.0.113.77 (INTERNET - ngoai MOI subnet interface cua
    # Template), rule duy nhat cover cap nay doi from=trust. Theo default route
    # trong Template thi zone that cua nguon la 'untrust' -> firewall THAT se chan.
    # TAI HIEN lo hong cu NGAY TRONG LAB bang zone_lookup KHONG co route.
    _zl_khong_route = build_zone_lookup(cfg.get('interface_subnets', []),
                                         cfg.get('zones', {}))
    _rb1_cu = classify_row(req.rows[16], ris, idx, zone_lookup=_zl_khong_route)
    check("Dong 17: B1 - TAI HIEN lo hong cu (khong route -> to xanh la FULL MATCH)",
          _rb1_cu.tier == "TOP1" and _compute_action_level(_rb1_cu) == "none"
          and _rb1_cu.zone_verdict == "unknown",
          f"tier={_rb1_cu.tier} action={_compute_action_level(_rb1_cu)} verdict={_rb1_cu.zone_verdict}")
    _rb1, _ab1 = row_of(16)
    check("Dong 17: B1 - co route (tu Template) thi suy duoc zone that va CHAN dung",
          _rb1.zone_verdict == "mismatch" and _ab1 != "none"
          and _rb1.tier not in ("TOP1", "TOP1.1", "TOP2", "TOP3"),
          f"tier={_rb1.tier} action={_ab1} verdict={_rb1.zone_verdict}")
    check("B1 - static route doc duoc tu Template (co tang <config> thua)",
          len(cfg.get('static_routes') or []) == 2,
          f"{len(cfg.get('static_routes') or [])} route")

    # --- B3 (10/2026): service cua rule deny KHONG resolve duoc ---
    # Dong 18: rule deny S00B-DENY-SVC-THIEU chan dung cap nay, nhung service cua no
    # tro tói mot service KHONG TON TAI -> resolve ra 0 PORT. Truoc ban va,
    # _ports_covered_for_shadow() coi 0 port la "khong cover port nao" nen tool IM
    # LANG - huong sai NGUY HIEM (admin tao rule moi duoi mot deny ma khong he biet).
    _rb3, _ab3 = row_of(17)
    check("Dong 18: B3 - deny co service khong resolve duoc VAN duoc canh bao",
          _rb3.blocking_deny_rule == "S00B-DENY-SVC-THIEU",
          f"deny={_rb3.blocking_deny_rule} tier={_rb3.tier}")
    check("Dong 18: B3 - va vi deny do o Shared-PRE nen ket luan la KHONG THE cap phep",
          _rb3.blocked_at_shared_pre, str(_rb3.blocking_deny_detail)[:70])

    # --- B4 (10/2026): deny o SHARED-POST - DOI XUNG voi P2, ket luan NGUOC DAU ---
    # Dong 19: moi rule cua Device Group deu duoc danh gia TRUOC Shared-POST, nen
    # them rule allow vao DG LA GIAI PHAP HOP LE. Neu gop chung nhan voi P2 thi
    # admin se tuong "bi chan boi Shared = khong lam duoc gi" roi di xin quyen
    # Panorama mot cach khong can thiet.
    _rb4, _ab4 = row_of(18)
    check("Dong 19: B4 - nhan dien deny nam o SHARED-POST",
          _rb4.blocking_deny_rule == "S98B-DENY-POST" and _rb4.blocked_at_shared_post,
          f"deny={_rb4.blocking_deny_rule} shared_post={_rb4.blocked_at_shared_post}")
    check("Dong 19: B4 - ket luan NGUOC DAU voi P2 (them rule vao DG la giai quyet duoc)",
          "KHONG can sua Shared" in (_rb4.blocking_deny_detail or "")
          and not _rb4.blocked_at_shared_pre,
          str(_rb4.blocking_deny_detail)[:80])
    _cb4 = "".join(getattr(b, "text", "") or ""
                   for b in build_name_rule_cell(_rb4)[0]).upper()
    check("Dong 19: B4 - nhan o Name Rule KHAC han P2",
          "SHARED-POST" in _cb4 and "KHONG THE CAP PHEP TU DEVICE GROUP NAY" not in _cb4,
          "nhan rieng cho Shared-POST")

    check("Dong 13: dung CANDIDATE config (khong bao TAO MOI)",
          results[12].used_candidate_rule is not None, str(results[12].used_candidate_rule))
    check("Dong 15: token rac duoc bao cao, khong nuot im lang",
          results[14].unparsed_source == ["KHONG-PHAI-IP"], str(results[14].unparsed_source))

    # --- B7 (10/2026): rule CUC BO tren chinh firewall (buoc 3 cua 0.3) ---
    # assemble_security_rules_panorama() co tham so local_rules va vi tri lap rap
    # dung tu lau, nhung GUI chua co duong nao nap vao nen no VINH VIEN la None ->
    # mot DENY cuc bo hoan toan VO HINH voi tool.
    from check_rule_pan.core import parse_local_firewall_rules
    from check_rule_pan.request import build_manual_request as _bmr
    _local_path = os.path.join(HERE, "fixtures", "firewall_local_config.xml")
    _local = parse_local_firewall_rules(_local_path)
    check("B7 - doc duoc rule cuc bo tu config cua CHINH firewall",
          [r.name for r in _local] == ["LOCAL-DENY-OPS"]
          and _local[0].vi_tri == "Local",
          f"{[(r.name, r.vi_tri) for r in _local]}")

    # Lap rap LAI co local_rules va kiem dung vi tri trong chuoi 0.3
    _cfg_local = load_pan_config_from_xml_string(xml, DG, SERIAL, local_rules=_local)
    _thu_tu_local = list(dict.fromkeys(r.vi_tri for r in _cfg_local["rules"]))
    check("B7 - rule cuc bo duoc lap rap DUNG vi tri (giua DG-PRE va DG-POST)",
          "Local" in _thu_tu_local
          and _thu_tu_local.index("Local") > _thu_tu_local.index("DG-HN-PRE")
          and _thu_tu_local.index("Local") < _thu_tu_local.index("DG-HN-POST"),
          " -> ".join(_thu_tu_local))

    # Va DENY cuc bo phai duoc canh bao - truoc B7 no vo hinh
    _ris_local = build_rule_infos(_cfg_local["rules"], GroupResolver(
        _cfg_local["address_objects"], _cfg_local["address_groups"]), ServiceResolver(
        _cfg_local["service_objects"], _cfg_local["service_groups"]))
    _row_local = _bmr([["10.10.55.11", "192.168.55.21", "8888"]]).rows[0]
    _r_local = classify_row(_row_local, _ris_local, idx, zone_lookup=zl)
    check("B7 - DENY cuc bo duoc canh bao (truoc B7 no VO HINH voi tool)",
          _r_local.blocking_deny_rule == "LOCAL-DENY-OPS",
          f"deny={_r_local.blocking_deny_rule}")
    # Doi chung: KHONG nap local_rules thi tool khong thay gi
    _r_khong = classify_row(_row_local, ris, idx, zone_lookup=zl)
    check("B7 - TAI HIEN: khong nap config firewall thi deny cuc bo VO HINH",
          _r_khong.blocking_deny_rule is None,
          f"deny={_r_khong.blocking_deny_rule}")

    out = os.path.join(OUT, "check_rule_ket_qua.xlsx")
    export_results(req.headers, results, out)
    import openpyxl
    ws = openpyxl.load_workbook(out)["Ket qua kiem tra"]
    hdr = [c.value for c in ws[1]][:5]
    check("Xuat Excel co cot 'Vi tri Rule' (0.3)",
          hdr == ["Name Rule", "Vi tri Rule", "Source", "Dest", "Port"], str(hdr))
    check("Cot 'Vi tri Rule' dien dung cho tung dong",
          ws.cell(row=2, column=2).value == "DG-HN-PRE"
          and ws.cell(row=6, column=2).value == "DG-DC-PRE",
          f"dong1={ws.cell(row=2, column=2).value} dong5={ws.cell(row=6, column=2).value}")
    legend = " ".join(str(c.value or "") for row in openpyxl.load_workbook(out)["Chu thich"].iter_rows() for c in row)
    check("Sheet Chu thich canh bao 2 TANG COMMIT cua Panorama (0.4)",
          "Commit and Push" in legend)
    return out


def tab_check_rule_live(host, key):
    section("TAB 1b - LIVE MATCH qua Panorama (0.5: BAT BUOC target=<serial>)")
    from check_rule_pan.core import fetch_panorama_config_xml, load_pan_config_from_xml_string
    from check_rule_pan.resolver import GroupResolver, ServiceResolver
    from check_rule_pan.index import build_device_index
    from check_rule_pan.matcher import build_rule_infos, classify_row
    from check_rule_pan.request import build_manual_request

    cfg = load_pan_config_from_xml_string(fetch_panorama_config_xml(host, key), DG, SERIAL)
    gr = GroupResolver(cfg["address_objects"], cfg["address_groups"])
    sr = ServiceResolver(cfg["service_objects"], cfg["service_groups"])
    idx = build_device_index(cfg, gr, sr)
    ris = build_rule_infos(cfg["rules"], gr, sr)

    row = build_manual_request([["10.10.10.11", "192.168.50.21", "1433"]]).rows[0]

    # 0.5: CO target -> Panorama forward xuong dataplane firewall
    r = classify_row(row, ris, idx, mgmt_ip=host, api_key=key, vsys="vsys1",
                     target_serial=SERIAL)
    check("0.5 - CO Target: Panorama forward -> firewall xac nhan rule that",
          r.tier == "LIVE_MATCH" and r.rule and r.rule.name == "H01-APP-TO-DB",
          f"tier={r.tier} rule={getattr(r.rule, 'name', None)}")

    # 0.5: THIEU target -> tu bo qua IM LANG, KHONG goi API
    calls = {"n": 0}
    import check_rule_pan.matcher as _m
    _orig = _m.fetch_security_policy_match
    def _spy(*a, **k):
        calls["n"] += 1
        return _orig(*a, **k)
    _m.fetch_security_policy_match = _spy
    try:
        r0 = classify_row(row, ris, idx, mgmt_ip=host, api_key=key, vsys="vsys1")
    finally:
        _m.fetch_security_policy_match = _orig
    check("0.5 - THIEU Target: bo qua LIVE_MATCH, KHONG goi API, roi xuong LOCAL",
          calls["n"] == 0 and r0.tier != "LIVE_MATCH" and r0.live_match_skip_reason is None,
          f"{calls['n']} API call, tier={r0.tier}")

    row2 = build_manual_request([["10.66.66.50", "192.168.50.21", "1433"]]).rows[0]
    r2 = classify_row(row2, ris, idx, mgmt_ip=host, api_key=key, vsys="vsys1",
                      target_serial=SERIAL)
    check("Combo bi DENY -> khong ket luan LIVE_MATCH, roi xuong LOCAL",
          r2.tier != "LIVE_MATCH", f"tier={r2.tier}")


# ======================================================================= TAB 2
def tab_check_ip(host, key):
    section(f"TAB 2 - CHECK IP  (pham vi = Shared + chuoi DG cua {DG})")
    import threading as _th
    from check_ip_pan.pan_api import (PanApiClient, get_devicegroup_list,
                                       resolve_devicegroup_chain)
    from check_ip_pan.core import parse_ip_list_input, build_run_plan, describe_run_plan
    from check_ip_pan.matching_engine import parse_log_file, run_search_pipeline, run_nat_export_pipeline
    from check_ip_pan.excel_export import (export_to_excel, export_nat_to_excel,
                                        export_combined_to_excel, RULE_SHEET_TITLE, NAT_SHEET_TITLE)

    client = PanApiClient(host, key, verify_ssl=False)
    check("API test_connection", client.test_connection() == "MOCK-FW-TESTLAB", client.test_connection())
    root = client.fetch_running_config()
    check("0.1 - doc danh sach Device Group",
          get_devicegroup_list(root) == ["DG-ALL", "DG-DC", "DG-HN"],
          str(get_devicegroup_list(root)))
    chain, _els = resolve_devicegroup_chain(root, DG)
    check("0.2 - chuoi thua ke cua DG dang chon",
          chain == ["shared", "DG-ALL", "DG-DC", "DG-HN"], " -> ".join(chain))
    tags = client.fetch_registered_ip()
    check("API doc registered-ip (Dynamic Address Group)", len(tags) == 2, f"{len(tags)} IP co tag")

    ip_text = open(os.path.join(HERE, "samples", "ip_list.txt"), encoding="utf-8").read()
    ips, bad = parse_ip_list_input(ip_text)
    check("Doc danh sach IP (co range day du / rut gon / co khoang trang)",
          len(ips) == 12 and not bad, f"{len(ips)} IP, {len(bad)} dong loi")

    plan, errs = build_run_plan(ip_text, "", False, "10.10.10.11", "xlsx", 30)
    check("Luat 1 nut 'Check / Export': nhap ca 2 -> chay CA HAI",
          plan["do_rule"] and plan["do_nat"] and not errs, describe_run_plan(plan))
    _p, e2 = build_run_plan("", "", False, "", "xlsx", 30)
    check("De trong het -> CHAN lai + noi ro can nhap gi", bool(e2), e2[0].splitlines()[0] if e2 else "")

    logs = parse_log_file(os.path.join(HERE, "samples", "traffic_log.xlsx"))
    check("Doc File Log doi chieu", len(logs) == 5, f"{len(logs)} dong traffic")

    rules, match, resolver, so, sg = run_search_pipeline(
        root, DG, ips, 20, log_func=lambda m: None,
        progress_cb=lambda *a: None, cancel_event=_th.Event(), registered_ip_tags=tags)
    check("Pipeline tra cuu Security Rule (ca Shared lan chuoi DG, pre+post)",
          len(rules) == 17 and len(match) > 0,
          f"{len(rules)} rule, {len(match)} IP co ket qua")
    vt = {r.name: r.vi_tri for r in rules}
    check("0.3 - moi rule co nhan 'Vi tri Rule' dung cho",
          vt.get("S00-BLOCK-BADNET") == "Shared-PRE"
          and vt.get("H01-APP-TO-DB") == "DG-HN-PRE"
          and vt.get("S99-ANY-ANY-ANY") == "Shared-POST",
          f"S00={vt.get('S00-BLOCK-BADNET')} H01={vt.get('H01-APP-TO-DB')}")
    check("Tim dung rule cho 10.10.10.11", any(True for _ in match.get("10.10.10.11", [])),
          f"{len(match.get('10.10.10.11', []))} tham chieu")

    nat_rules, nat_res, nso, nsg = run_nat_export_pipeline(
        root, DG, log_func=lambda m: None, registered_ip_tags=tags)
    check("Pipeline doc bang NAT theo DG", len(nat_rules) == 2, f"{len(nat_rules)} NAT rule")
    nvt = {r.name: r.vi_tri for r in nat_rules}
    check("0.3 - NAT cung co nhan 'Vi tri Rule' (pre/post-rulebase cua DG)",
          nvt.get("NAT01-OUTBOUND-HIDE") == "DG-HN-PRE"
          and nvt.get("NAT02-INBOUND-DNAT") == "DG-HN-POST", str(nvt))

    p_rule = export_to_excel(OUT, rules, match, resolver, so, sg, log_entries=logs,
                              log_func=lambda m: None)
    import openpyxl as _ox
    _h = [c.value for c in _ox.load_workbook(p_rule).worksheets[0][1]][:5]
    check("Xuat Excel Rule: cot 'Vi tri Rule' + 3 cot Log doi chieu",
          _h == ["Name Rule", "Vi tri Rule", "Log Source IP", "Log Dest IP", "Log Port"], str(_h))
    p_nat = export_nat_to_excel(OUT, nat_rules, nat_res, nso, nsg, log_func=lambda m: None)
    check("Xuat Excel NAT", os.path.exists(p_nat), os.path.basename(p_nat))

    p_comb = export_combined_to_excel(
        OUT, rules=rules, match_results=match, resolver=resolver,
        service_objects=so, service_groups=sg,
        nat_rules=nat_rules, nat_resolver=nat_res,
        nat_service_objects=nso, nat_service_groups=nsg,
        log_entries=logs, log_func=lambda m: None)
    import openpyxl
    wb = openpyxl.load_workbook(p_comb)
    check("Gop 1 FILE 2 SHEET khi chay ca Rule lan NAT",
          wb.sheetnames == [RULE_SHEET_TITLE, NAT_SHEET_TITLE], str(wb.sheetnames))
    return p_comb


# ======================================================================= TAB 3
def tab_check_audit(host, key):
    section(f"TAB 3 - CHECK AUDIT  (Device Group={DG}, Target={SERIAL})")
    from check_audit_pan.pan_api import (PanApiClient, parse_devicegroup_objects,
                                          assemble_security_rules, fetch_rule_hit_count,
                                          disable_rule_candidate, delete_rule_candidate)
    from check_audit_pan.resolver import GroupResolver, ServiceResolver
    from check_audit_pan.detectors import find_shadow_rules, find_disabled_rules, find_overly_permissive_rules
    from check_audit_pan.excel_io import export_findings

    client = PanApiClient(host, key, verify_ssl=False)
    root = client.fetch_running_config()
    objs = parse_devicegroup_objects(root, DG)
    rules = assemble_security_rules(root, DG)
    check("Doc rulebase cho audit (lap rap theo 0.3)", len(rules) == 17, f"{len(rules)} rule")

    gr = GroupResolver(objs["address_objects"], objs["address_groups"])
    sr = ServiceResolver(objs["service_objects"], objs["service_groups"])
    import logging as _lg
    # --- B2 (10/2026): rule Shared ap cho MOI firewall Panorama quan ly ---
    # DG-HN co 2 firewall, nhung Panorama con quan ly cac con khac. Rule o
    # Shared-PRE/Shared-POST ap cho TAT CA, nen "0 hit tren moi con trong DG" VAN
    # CHUA DU bang chung -> co the de xuat xoa rule Shared dang chay o DG khac.
    from check_audit_pan.pan_api import (get_panorama_device_serials,
                                          devicegroup_device_serials,
                                          fetch_rule_hit_count_multi)
    # `_serials` (cac firewall cua DG) duoc tinh o khoi P1 BEN DUOI, nen o day tinh
    # rieng - khong dua vao thu tu khoi trong file.
    _serials = devicegroup_device_serials(root, DG)
    _estate = get_panorama_device_serials(host, key, log_func=lambda m: None)
    check("B2 - liet ke duoc MOI firewall Panorama quan ly",
          len(_estate) >= 2 and SERIAL in _estate, str(_estate))
    _ngoai = [x for x in _estate if x not in _serials]
    _hits_estate = fetch_rule_hit_count_multi(host, key, _serials,
                                               log_func=lambda m: None,
                                               extra_shared_serials=_ngoai)
    _st_shared = _hits_estate.get(("S10-SHADOWED-NOHIT", "Shared-POST")) \
        or next((v for k, v in _hits_estate.items()
                 if str(k[1]).startswith("Shared-")), None)
    check("B2 - rule o pham vi Shared co bang chung tu DU cac firewall can kiem",
          _st_shared is not None
          and set(_st_shared.devices_seen) >= set(_serials),
          str(_st_shared and _st_shared.evidence_scope)[:90])
    _st_dg = next((v for k, v in _hits_estate.items()
                   if str(k[1]).startswith("DG-")), None)
    check("B2 - rule trong Device Group KHONG doi hoi bang chung tu firewall ngoai DG",
          _st_dg is not None and _st_dg.hits_known
          and all(x in _serials for x in _st_dg.devices_seen),
          str(_st_dg and _st_dg.evidence_scope)[:90])

    # --- P1 (10/2026): UNUSED phai 0 hit tren MOI firewall cua Device Group ---
    # So sanh TRUC TIEP 2 cach doc tren cung 1 rule (H04-DUNG-O-FW2-NOHIT1):
    #   - doc rieng serial 001122334455 -> 0 hit -> ket luan UNUSED (CACH CU, NGUY HIEM)
    #   - doc ca 2 firewall cua DG-HN    -> co hit -> KHONG phai UNUSED (CACH DUNG)
    # Rule nay nam o DG-HN-PRE va DG-HN co 2 firewall, nen "UNUSED" theo cach cu
    # dong nghia voi de xuat XOA mot rule dang chay tren firewall thu 2.
    import datetime as _dt
    from check_audit_pan.detectors import _find_unused_rules as _unused
    from check_audit_pan.pan_api import (fetch_rule_hit_count_multi,
                                          devicegroup_device_serials)

    _serials = devicegroup_device_serials(root, DG)
    check("P1 - liet ke dung cac firewall cua Device Group (khong lay nham ban readonly)",
          len(_serials) == 2 and SERIAL in _serials,
          f"{DG} -> {_serials}")

    _rule_h04 = [r for r in rules if r.name == "H04-DUNG-O-FW2-NOHIT1"]
    check("P1 - fixture co rule dung o firewall thu 2", len(_rule_h04) == 1,
          str([r.name for r in rules if "H04" in r.name]))

    _now = _dt.datetime.now()
    # (a) CACH CU: chi doc 1 serial
    _hits_one = fetch_rule_hit_count(host, key, SERIAL, log_func=lambda m: None)
    _unused_one = _unused(_rule_h04, _hits_one, _now)
    check("P1 - TAI HIEN rui ro: doc 1 serial thi rule nay bi ket luan UNUSED",
          bool(_unused_one),
          f"hit tren {SERIAL} = {(_hits_one.get(('H04-DUNG-O-FW2-NOHIT1','DG-HN-PRE')) or object()).__dict__.get('hit_count', '?') if _hits_one.get(('H04-DUNG-O-FW2-NOHIT1','DG-HN-PRE')) else 'khong co'}")

    # (b) CACH DUNG: doc ca Device Group
    _hits_all = fetch_rule_hit_count_multi(host, key, _serials, log_func=lambda m: None)
    _st = _hits_all.get(("H04-DUNG-O-FW2-NOHIT1", "DG-HN-PRE"))
    check("P1 - doc ca DG thi thay rule DANG DUOC DUNG (hit > 0)",
          _st is not None and (_st.hit_count or 0) > 0,
          f"hit gop = {_st and _st.hit_count}")
    check("P1 - va KHONG con bi ket luan UNUSED (khong de xuat xoa rule dang chay)",
          _unused(_rule_h04, _hits_all, _now) == {}, "")
    check("P1 - bang chung ghi ro da kiem nhung firewall nao",
          _st is not None and set(_st.devices_seen) == set(_serials)
          and _st.hits_known and _st.evidence_scope,
          str(_st and _st.evidence_scope))

    # (c) THIEU bang chung: 1 firewall khong doc duoc -> "khong biet" KHAC "bang 0"
    _cli_p1 = PanApiClient(host, key, verify_ssl=False)
    _orig_p1 = _cli_p1._request
    def _reject_fw2(params, *a, **k):
        if params.get("target") == "009988776655":
            raise RuntimeError("mo phong: firewall thu 2 khong phan hoi")
        return _orig_p1(params, *a, **k)
    _cli_p1._request = _reject_fw2
    _lg.getLogger("CheckAudit").setLevel(_lg.CRITICAL)
    try:
        _hits_part = _cli_p1.fetch_rule_hit_count_multi(_serials, log_func=lambda m: None)
    finally:
        _lg.getLogger("CheckAudit").setLevel(_lg.WARNING)
    _rule_dead = [r for r in rules if r.name == "H90-SHADOWED-NOHIT"]
    _st_part = _hits_part.get(("H90-SHADOWED-NOHIT", "DG-HN-POST"))
    check("P1 - firewall khong doc duoc -> hits_known=False (khong biet KHAC bang 0)",
          _st_part is not None and not _st_part.hits_known
          and "009988776655" in _st_part.devices_missing,
          str(_st_part and _st_part.evidence_scope))
    check("P1 - thieu bang chung thi KHONG ket luan UNUSED du phan doc duoc la 0 hit",
          _unused(_rule_dead, _hits_part, _now) == {}, "")

    hits_notarget = fetch_rule_hit_count(host, key, "", log_func=lambda m: None)
    check("0.7 - THIEU Target: hit-count tra {} (nhanh UNUSED tu tat)",
          hits_notarget == {}, str(hits_notarget))
    hits = fetch_rule_hit_count(host, key, SERIAL, log_func=lambda m: None)
    check("0.7 - CO Target: doc duoc Rule Hit Count (Policy Optimizer)",
          len(hits) >= 13, f"{len(hits)} rule co hit-count")
    zero = [n for n, s in hits.items() if (s.hit_count or 0) == 0]
    check("Co rule 0 hit de kiem thu nhanh UNUSED", len(zero) >= 2, f"{len(zero)} rule 0 hit")
    check("D9 - hit-count khoa theo (ten rule, Vi tri Rule), suy tu chinh response",
          ("S00-BLOCK-BADNET", "Shared-PRE") in hits
          and ("S99-ANY-ANY-ANY", "Shared-POST") in hits
          and ("H01-APP-TO-DB", "DG-HN-PRE") in hits
          and ("H90-SHADOWED-NOHIT", "DG-HN-POST") in hits,
          "du ca 4 pham vi Shared-PRE/POST + DG-PRE/POST")
    # Panorama that co the KHONG ho tro node <shared> cho rule-hit-count -> phai
    # suy giam mem: rule trong Device Group van doc duoc, khong raise, khong rong.
    _cli = PanApiClient(host, key, verify_ssl=False)
    _orig_req = _cli._request
    def _reject_shared(params, *a, **k):
        if "<shared>" in (params.get("cmd") or ""):
            raise RuntimeError("mo phong: Panorama tu choi node <shared>")
        return _orig_req(params, *a, **k)
    _cli._request = _reject_shared
    _warn = []
    _lg.getLogger("CheckAudit").setLevel(_lg.CRITICAL)  # canh bao nay la CO CHU DICH
    try:
        _degraded = _cli.fetch_rule_hit_count(SERIAL, log_func=_warn.append)
    finally:
        _lg.getLogger("CheckAudit").setLevel(_lg.NOTSET)
    check("0.7 - Panorama tu choi pham vi <shared>: suy giam mem, DG van co hit-count",
          len(_degraded) >= 11
          and not any(k[1].startswith("Shared") for k in _degraded)
          and any("CANH BAO" in m for m in _warn),
          f"{len(_degraded)} rule DG + {sum('CANH BAO' in m for m in _warn)} canh bao")

    now = datetime.datetime.now()
    f_shadow = find_shadow_rules(rules, gr, sr, hits, now)
    f_disabled = find_disabled_rules(rules)
    f_overly = find_overly_permissive_rules(rules)
    names_shadow = {f.rule_name for f in f_shadow}

    check("Phat hien DISABLED rule", {f.rule_name for f in f_disabled} == {"D90-DISABLED-NOHIT"},
          str({f.rule_name for f in f_disabled}))
    check("Phat hien OVERLY PERMISSIVE (any/any/any o Shared-POST)",
          "S99-ANY-ANY-ANY" in {f.rule_name for f in f_overly},
          str({f.rule_name for f in f_overly}))
    check("Phat hien SHADOW that (rule DG-HN-POST bi rule DG-HN-PRE che)",
          "H90-SHADOWED-NOHIT" in names_shadow, f"{len(f_shadow)} finding shadow")
    check("FIX ZONE: KHONG bao shadow gia khi R2 dung zone 'any'",
          "D03B-ZONE-ANY-OK" not in names_shadow,
          "D03B (zone any) khong bi D03 (chi trust->untrust) che" if
          "D03B-ZONE-ANY-OK" not in names_shadow else "VAN BAO SHADOW GIA!")

    findings = f_shadow + f_disabled + f_overly
    p = export_findings(findings, OUT, "xlsx")
    import openpyxl as _ox2
    _ws = _ox2.load_workbook(p).worksheets[0]
    _hdr = [c.value for c in _ws[1]]
    check("Xuat file audit co cot 'Vi tri Rule' (15 cot)",
          len(_hdr) == 15 and "Vi tri Rule" in _hdr, f"{len(_hdr)} cot")

    # D6: cac lenh ghi chi ap len CANDIDATE, KHONG BAO GIO commit
    jpath = os.path.join(HERE, ".journal.log")
    before = os.path.getsize(jpath) if os.path.exists(jpath) else 0
    disable_rule_candidate(host, key, DG, "post", "H90-SHADOWED-NOHIT")
    delete_rule_candidate(host, key, DG, "post", "D90-DISABLED-NOHIT")
    after_txt = open(jpath, encoding="utf-8").read()[before:] if os.path.exists(jpath) else ""
    check("0.8 - Import Shadow(hit=0) -> DISABLE, xpath dung post-rulebase cua DG",
          "H90-SHADOWED-NOHIT" in after_txt and "post-rulebase" in after_txt
          and f"device-group/entry[@name='{DG}']" in after_txt,
          [l for l in after_txt.splitlines() if "H90" in l][:1])
    check("0.8 - Import Disabled -> DELETE, xpath dung DG",
          "action=delete" in after_txt and "D90-DISABLED-NOHIT" in after_txt)
    check("D6/0.4: TUYET DOI khong goi commit (Panorama con can Commit and Push)",
          "commit" not in after_txt.lower())
    return p


# ====================================================================== runner
@contextlib.contextmanager
def auto_mock(port):
    proc = subprocess.Popen([sys.executable, os.path.join(HERE, "mock_panorama_api.py"),
                             "--port", str(port), "--quiet"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.STDOUT)
    try:
        import urllib.request, ssl as _ssl
        ctx = _ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = _ssl.CERT_NONE
        for _ in range(60):
            try:
                urllib.request.urlopen(f"https://127.0.0.1:{port}/api/?key=x&type=op&cmd=<show><system><info></info></system></show>",
                                       context=ctx, timeout=1).read()
                break
            except Exception:
                time.sleep(0.25)
        yield f"127.0.0.1:{port}"
    finally:
        proc.terminate()
        with contextlib.suppress(Exception):
            proc.wait(timeout=5)


def main():
    ap = argparse.ArgumentParser(description="Chay thu toan bo tinh nang qua Panorama gia lap")
    ap.add_argument("--host", default=None, help="vd 127.0.0.1:8443 (mock da chay san)")
    ap.add_argument("--key", default="testlab")
    ap.add_argument("--auto", action="store_true", help="tu bat/tat mock (1 lenh duy nhat)")
    ap.add_argument("--port", type=int, default=18443, help="port khi dung --auto")
    args = ap.parse_args()

    if not args.host and not args.auto:
        ap.error("Can --host 127.0.0.1:8443 (mock da chay), hoac --auto de tu bat mock.")

    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT, exist_ok=True)
    jpath = os.path.join(HERE, ".journal.log")
    if os.path.exists(jpath):
        os.remove(jpath)

    print("=" * 78)
    print("  NGHIEM THU TOAN BO TINH NANG - PaloPANORAMA (qua Panorama gia lap)")
    print("=" * 78)

    ctx = auto_mock(args.port) if args.auto else contextlib.nullcontext(args.host)
    with ctx as host:
        print(f"  Thiet bi: https://{host}/api/   (API Key: {args.key})")
        tab_check_rule(host, args.key)
        tab_check_rule_live(host, args.key)
        tab_check_ip(host, args.key)
        tab_check_audit(host, args.key)

    n_pass = sum(1 for okc, _, _ in RESULTS if okc)
    n_fail = len(RESULTS) - n_pass
    print()
    print("=" * 78)
    print(f"  TONG KET: {n_pass}/{len(RESULTS)} PASS" + (f", {n_fail} FAIL" if n_fail else ""))
    print(f"  File ket qua: {OUT}")
    print("=" * 78)
    if n_fail:
        print("\n  CAC MUC FAIL:")
        for okc, name, detail in RESULTS:
            if not okc:
                print(f"    - {name}  ({detail})")
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
