# -*- coding: utf-8 -*-
"""
E2E tab CHECK IP ban PANORAMA: chay THAT pipeline tren cau hinh Panorama gia lap
(Shared + 3 tang Device Group, pre/post-rulebase) roi xuat Excel va doc lai.

Kiem chung:
  - Pham vi tra cuu = Shared + CHUOI DG cha-con (khong phai 1 vsys)
  - Cot "Vi tri Rule" ghi dung noi rule nam (III.7)
  - 1 luot chay ca Rule lan NAT -> DUNG 1 file 2 sheet
Khong phu thuoc tkinter, khong goi mang that.
"""

import os
import shutil
import tempfile
import threading
import unittest
import xml.etree.ElementTree as ET

import openpyxl

from tests.panorama_fixture import PANORAMA_XML
from check_ip_pan.pan_api import get_devicegroup_list, resolve_devicegroup_chain
from check_ip_pan.matching_engine import run_search_pipeline, run_nat_export_pipeline
from check_ip_pan.excel_export import (export_to_excel, export_nat_to_excel,
                                        export_combined_to_excel,
                                        RULE_SHEET_TITLE, NAT_SHEET_TITLE)


def _root():
    return ET.fromstring(PANORAMA_XML)


class TestCheckIpPanoramaScope(unittest.TestCase):
    def test_danh_sach_device_group(self):
        self.assertEqual(get_devicegroup_list(_root()), ["DG-ALL", "DG-DC", "DG-HN"])

    def test_chuoi_thua_ke(self):
        names, _ = resolve_devicegroup_chain(_root(), "DG-HN")
        self.assertEqual(names, ["shared", "DG-ALL", "DG-DC", "DG-HN"])


class TestCheckIpPanoramaPipeline(unittest.TestCase):
    def setUp(self):
        self.out = tempfile.mkdtemp(prefix="checkip_pan_")
        self.logs = []

    def tearDown(self):
        shutil.rmtree(self.out, ignore_errors=True)

    def _run(self, dg, ips):
        return run_search_pipeline(
            _root(), dg, ips, 20,
            log_func=self.logs.append,
            progress_cb=lambda *a: None,
            cancel_event=threading.Event(),
            registered_ip_tags={})

    def test_tim_thay_rule_o_moi_tang_cua_chuoi_dg(self):
        """1 lan tra cuu phai quet CA Shared, DG to tien, DG cha va DG hien tai."""
        rules, match, resolver, so, sg = self._run("DG-HN", ["10.7.7.7", "10.20.20.20"])
        vi_tri = {r.name: r.vi_tri for r in rules}
        self.assertEqual(vi_tri["SHARED-PRE-DENY"], "Shared-PRE")
        self.assertEqual(vi_tri["DGALL-PRE-ALLOW"], "DG-ALL-PRE")
        self.assertEqual(vi_tri["DGHN-POST"], "DG-HN-POST")
        self.assertEqual(vi_tri["SHARED-POST-CATCHALL"], "Shared-POST")
        # .7.7 la object HN-HOST cua DG hien tai; .20.20 nam o rule cua DG to tien
        self.assertIn("10.7.7.7", match)
        self.assertIn("10.20.20.20", match)

    def test_dg_khac_pham_vi_khac(self):
        rules_all, _, _, _, _ = self._run("DG-ALL", ["10.7.7.7"])
        rules_hn, _, _, _, _ = self._run("DG-HN", ["10.7.7.7"])
        self.assertLess(len(rules_all), len(rules_hn))
        self.assertNotIn("DGHN-PRE", [r.name for r in rules_all])

    def test_export_co_cot_vi_tri_rule(self):
        rules, match, resolver, so, sg = self._run("DG-HN", ["10.7.7.7"])
        path = export_to_excel(self.out, rules, match, resolver, so, sg,
                                log_func=self.logs.append)
        ws = openpyxl.load_workbook(path).worksheets[0]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[0], "Name Rule")
        self.assertEqual(headers[1], "Vi tri Rule")
        col = {h: i for i, h in enumerate(headers)}
        found = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[col["Name Rule"]]:
                found[row[col["Name Rule"]]] = row[col["Vi tri Rule"]]
        self.assertTrue(found, "Khong xuat duoc rule nao")
        for name, vt in found.items():
            self.assertTrue(vt, f"rule {name} thieu Vi tri Rule")

    def test_nat_pipeline_theo_dg(self):
        nat_rules, resolver, so, sg = run_nat_export_pipeline(
            _root(), "DG-HN", log_func=self.logs.append, registered_ip_tags={})
        self.assertIsInstance(nat_rules, list)   # fixture khong co NAT -> rong, khong raise

    def test_gop_1_file_2_sheet(self):
        rules, match, resolver, so, sg = self._run("DG-HN", ["10.7.7.7"])
        nat_rules, nat_resolver, nso, nsg = run_nat_export_pipeline(
            _root(), "DG-HN", log_func=self.logs.append, registered_ip_tags={})
        path = export_combined_to_excel(
            self.out, rules=rules, match_results=match, resolver=resolver,
            service_objects=so, service_groups=sg,
            nat_rules=nat_rules, nat_resolver=nat_resolver,
            nat_service_objects=nso, nat_service_groups=nsg,
            log_func=self.logs.append)
        self.assertTrue(os.path.basename(path).startswith("rule_nat_"))
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, [RULE_SHEET_TITLE, NAT_SHEET_TITLE])
        self.assertEqual([c.value for c in wb[RULE_SHEET_TITLE][1]][:2],
                         ["Name Rule", "Vi tri Rule"])


if __name__ == "__main__":
    unittest.main()
