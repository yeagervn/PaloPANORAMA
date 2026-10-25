# -*- coding: utf-8 -*-
"""
E2E GUI tab CHECK IP: lai THAT giao dien (go vao o nhap, bam nut) tren
FIREWALL GIA LAP cua testlab, roi doi chieu file ket qua.

Vi sao can: PHAN 0.2 cua tai lieu ghi "GUI that (ve widget, bam nut): CHUA
kiem chung". testlab/run_all_features.py goi ham pipeline truc tiep - KHONG
di qua tang GUI, nen khong phu: nut co bat/tat dung luc khong, worker co day
duoc ket qua ve main thread khong, UI co bi ket o trang thai busy khi chay
xong / khi bam Cancel khong.

Tu skip neu may khong co python3-tk hoac khong co DISPLAY (dung xvfb-run).
"""

import os
import shutil
import sys
import tempfile
import time
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

MOCK_PORT = 18477          # khac cong mac dinh 8443 de khong dung lab dang chay


def _tk_ready():
    try:
        import tkinter as tk
    except Exception as exc:
        return None, f"khong co tkinter ({exc})"
    if sys.platform.startswith("linux") and not os.environ.get("DISPLAY"):
        return None, "khong co DISPLAY (chay qua xvfb-run)"
    try:
        r = tk.Tk()
        r.destroy()
    except Exception as exc:
        return None, f"khong mo duoc Tk: {exc}"
    return tk, ""


class _NoDialogs:
    """Chan moi hop thoai - test khong duoc dung cho nguoi bam OK."""

    def __init__(self, gui_mod):
        self.mod = gui_mod
        self.saved = {}
        self.calls = []

    def __enter__(self):
        for name, ret in (("showinfo", None), ("showwarning", None),
                          ("showerror", None), ("askyesno", True)):
            self.saved[name] = getattr(self.mod.messagebox, name)
            setattr(self.mod.messagebox, name,
                    (lambda *a, _n=name, _r=ret, **k: (self.calls.append((_n, a)), _r)[1]))
        return self

    def __exit__(self, *exc):
        for name, fn in self.saved.items():
            setattr(self.mod.messagebox, name, fn)
        return False


class TestCheckIpGuiEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        tk, why = _tk_ready()
        if tk is None:
            raise unittest.SkipTest(why)
        cls.tk = tk

        from testlab.run_all_features import auto_mock

        cls._mock_ctx = auto_mock(MOCK_PORT)
        cls.host = cls._mock_ctx.__enter__()

        from check_ip_pan import core as ip_core
        cls.out_dir = ip_core.OUTPUT_DIR

    @classmethod
    def tearDownClass(cls):
        if hasattr(cls, "_mock_ctx"):
            cls._mock_ctx.__exit__(None, None, None)

    # ------------------------------------------------------------- tien ich
    def setUp(self):
        # Test KHONG duoc de lai rac trong repo: ghi nho thu muc output truoc
        # khi chay, ket thuc thi xoa dung nhung file do test sinh ra.
        self._before = set(os.listdir(self.out_dir)) if os.path.isdir(self.out_dir) else set()
        self.addCleanup(self._cleanup_output)

    def _cleanup_output(self):
        if not os.path.isdir(self.out_dir):
            return
        for name in set(os.listdir(self.out_dir)) - self._before:
            path = os.path.join(self.out_dir, name)
            try:
                os.remove(path) if os.path.isfile(path) else shutil.rmtree(path)
            except OSError:
                pass

    def _pump(self, root, predicate, timeout=90, label=""):
        end = time.time() + timeout
        while time.time() < end:
            root.update()
            if predicate():
                return True
            time.sleep(0.05)
        self.fail(f"Qua {timeout}s van chua: {label}")

    def _new_app(self):
        from check_ip_pan.gui import App

        root = self.tk.Tk()
        root.withdraw()
        app = App(root)
        app.pack(fill="both", expand=True)
        root.update_idletasks()
        self.addCleanup(self._teardown_tk, app, root)
        return root, app

    @staticmethod
    def _teardown_tk(app, root):
        """Dong cua so VA buong het tham chieu Tk ngay tren main thread.

        Khong lam vay thi cac tk.Variable cua app bi GC sau khi interpreter Tcl
        da bien mat -> Variable.__del__ goi vao Tcl tu thread GC ->
        "Tcl_AsyncDelete: async handler deleted by the wrong thread" lam ABORT
        ca tien trinh test (chi lo ra khi chay CA bo test, khong lo khi chay
        rieng 1 file).
        """
        import gc

        import threading

        try:
            app.shutdown()          # bat cancel_event -> worker se dung som
        except Exception:
            pass

        # CHO worker thread ket thuc TRUOC khi buong tham chieu. Neu xoa
        # __dict__ trong luc worker con chay, no se nem AttributeError khi goi
        # self._log() - loi cua TEST chu khong phai cua tool.
        deadline = time.time() + 5
        while time.time() < deadline:
            alive = [t for t in threading.enumerate()
                     if t is not threading.main_thread() and t.is_alive()]
            if not alive:
                break
            time.sleep(0.05)

        try:
            root.destroy()
        except Exception:
            pass
        # Buong het tham chieu Tk ngay tren main thread roi gom rac: tk.Variable
        # bi GC sau khi interpreter Tcl da bien mat se lam Tcl abort ca tien
        # trinh ("Tcl_AsyncDelete: async handler deleted by the wrong thread").
        app.__dict__.clear()
        gc.collect()

    def _connect(self, root, app):
        app.ent_host.delete(0, "end")
        app.ent_host.insert(0, self.host)
        app.ent_apikey.delete(0, "end")
        app.ent_apikey.insert(0, "testlab")
        app.var_verify_ssl.set(False)

        self.assertEqual(str(app.btn_run.cget("state")), "disabled",
                         "chua Connect ma nut Check/Export da bat")
        app._on_connect()
        # Doi MAIN THREAD xu ly xong message 'done_connect' (chinh no moi do
        # danh sach VSYS va mo khoa nut) - khong doi co config_root, vi bien do
        # duoc worker thread gan TRUOC khi main thread kip cap nhat giao dien.
        self._pump(root, lambda: str(app.btn_run.cget("state")) == "normal",
                   label="ket noi xong va mo khoa UI")
        return app

    # ---------------------------------------------------------------- tests
    def test_connect_then_run_security_and_nat(self):
        import check_ip_pan.gui as gui_mod

        root, app = self._new_app()
        with _NoDialogs(gui_mod) as dialogs:
            self._connect(root, app)

            dgs = app.cmb_dg.cget("values")
            self.assertTrue(dgs, "khong lay duoc danh sach Device-Group")
            # Chon dung DG ma testlab chot so lieu (DG-HN co 2 NAT rule).
            # DG mac dinh la DG-ALL - KHONG co NAT rule nao, nen sheet NAT
            # rong o DG do la DUNG, khong phai loi.
            self.assertIn("DG-HN", dgs)
            app.cmb_dg.set("DG-HN")
            self.assertEqual(str(app.btn_run.cget("state")), "normal",
                             "Connect xong ma nut Check/Export van khoa")

            with open(os.path.join(ROOT, "testlab", "samples", "ip_list.txt"),
                      encoding="utf-8") as f:
                ips = f.read()
            app.txt_ips.delete("1.0", "end")
            app.txt_ips.insert("1.0", ips)
            app.var_export_nat.set("Yes")
            app._on_toggle_nat()

            before = set(self._before)
            app._on_open_output = lambda *a, **k: None      # khong mo file manager

            app._on_run_all()
            self.assertEqual(str(app.btn_run.cget("state")), "disabled",
                             "dang chay ma nut Check/Export van bam duoc")
            self._pump(root, lambda: str(app.btn_run.cget("state")) == "normal",
                       label="chay xong va mo khoa lai UI")

        new_files = sorted(set(os.listdir(self.out_dir)) - before)
        self.assertTrue(new_files, "chay xong nhung khong sinh file ket qua nao")
        xlsx = [f for f in new_files if f.endswith(".xlsx")]
        self.assertTrue(xlsx, f"khong co file .xlsx nao trong {new_files}")

        # Muc 2 + muc 4 cung chay -> dung 1 file gom 2 sheet (luat III.7)
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(self.out_dir, xlsx[0]))
        self.assertEqual(wb.sheetnames, ["Security Rules", "NAT Rules"],
                         f"sai cau truc sheet: {wb.sheetnames}")
        self.assertGreater(wb["Security Rules"].max_row, 1, "sheet Security Rules rong")
        self.assertGreater(wb["NAT Rules"].max_row, 1, "sheet NAT Rules rong")
        wb.close()
        self.assertTrue(any(n == "showinfo" for n, _ in dialogs.calls),
                        "chay xong khong bao cho nguoi dung biet file o dau")

    def test_run_with_empty_input_does_not_lock_ui(self):
        """De trong het roi bam Run: phai bao va mo khoa UI, khong treo."""
        import check_ip_pan.gui as gui_mod

        root, app = self._new_app()
        with _NoDialogs(gui_mod) as dialogs:
            self._connect(root, app)
            app.txt_ips.delete("1.0", "end")
            app.ent_logfile.delete(0, "end")
            app.var_export_nat.set("No")
            app.txt_nat_ips.delete("1.0", "end")

            app._on_run_all()
            root.update()
            self._pump(root, lambda: str(app.btn_run.cget("state")) == "normal",
                       timeout=30, label="UI duoc mo khoa lai")
        self.assertTrue(dialogs.calls, "khong canh bao gi khi khong nhap gi ca")

    def test_cancel_midway_releases_ui(self):
        """Bam Cancel giua chung: UI phai tro lai binh thuong, khong ket."""
        import check_ip_pan.gui as gui_mod

        root, app = self._new_app()
        with _NoDialogs(gui_mod):
            self._connect(root, app)
            app.txt_ips.delete("1.0", "end")
            app.txt_ips.insert("1.0", "10.1.1.1\n10.1.1.2\n")
            app.var_export_nat.set("Yes")
            app._on_toggle_nat()
            app._on_open_output = lambda *a, **k: None

            app._on_run_all()
            root.update()
            app._on_cancel()
            self.assertTrue(app.cancel_event.is_set())
            self._pump(root, lambda: str(app.btn_run.cget("state")) == "normal",
                       label="UI mo khoa lai sau khi Cancel")


if __name__ == "__main__":
    unittest.main(verbosity=2)
