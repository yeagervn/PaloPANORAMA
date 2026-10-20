# -*- coding: utf-8 -*-
"""
E2E tab CHECK RULE ban PANORAMA: chay THAT tu cau hinh Panorama gia lap -> classify
-> xuat Excel -> doc lai file. Kiem chung ca luat nghiep vu (giu nguyen 100% ban
Firewall) lan phan ha tang Panorama (cot "Vi tri Rule", pham vi DG-chain).
"""

import os
import shutil
import tempfile
import unittest

import openpyxl

from tests.panorama_fixture import PANORAMA_XML
from check_rule_pan.core import load_pan_config_from_xml_string
from check_rule_pan.resolver import GroupResolver, ServiceResolver
from check_rule_pan.index import build_device_index, build_zone_lookup
from check_rule_pan.matcher import build_rule_infos, classify_row
from check_rule_pan.request import build_manual_request
from check_rule_pan.exporter import export_results, _compute_action_level


def _ctx(dg="DG-HN", serial=None):
    cfg = load_pan_config_from_xml_string(PANORAMA_XML, dg, serial)
    gr = GroupResolver(cfg["address_objects"], cfg["address_groups"])
    sr = ServiceResolver(cfg["service_objects"], cfg["service_groups"])
    idx = build_device_index(cfg, gr, sr)
    ris = build_rule_infos(cfg["rules"], gr, sr)
    zl = build_zone_lookup(cfg["interface_subnets"], cfg["zones"])
    return cfg, idx, ris, zl


def _one(src, dst, port, dg="DG-HN", serial=None):
    cfg, idx, ris, zl = _ctx(dg, serial)
    row = build_manual_request([[src, dst, port]]).rows[0]
    return classify_row(row, ris, idx, zone_lookup=zl)


class TestTierTrenPanorama(unittest.TestCase):
    def test_top1_qua_object_cua_dg_hien_tai(self):
        r = _one("10.7.7.7", "192.168.100.10", "443")
        self.assertEqual(r.tier, "TOP1")
        self.assertEqual(r.rule.name, "DGHN-PRE")
        self.assertEqual(r.rule.vi_tri, "DG-HN-PRE")
        self.assertEqual(_compute_action_level(r), "none")

    def test_top1_qua_object_THUA_KE_tu_dg_cha(self):
        """DC-RANGE dinh nghia o DG-DC, rule DGDC-PRE cung o DG-DC - DG-HN thua ke ca 2."""
        r = _one("10.5.5.100-10.5.5.110", "192.168.100.10", "443")
        self.assertEqual(r.tier, "TOP1")
        self.assertEqual(r.rule.vi_tri, "DG-DC-PRE")

    def test_range_con_nam_trong_range_object_thua_ke(self):
        r = _one("10.5.5.102-10.5.5.105", "192.168.100.10", "443")
        self.assertEqual(r.tier, "TOP1")
        self.assertEqual(r.source_report[0].status, "in_rule")

    def test_rule_cua_dg_to_tien_van_duoc_ap_dung(self):
        r = _one("10.20.20.20", "192.168.100.10", "443")
        self.assertEqual(r.rule.vi_tri, "DG-ALL-PRE")

    def test_dg_khac_khong_thay_rule_cua_dg_em(self):
        """Dung o DG-ALL thi KHONG duoc thay rule/object rieng cua DG-HN."""
        r = _one("10.7.7.7", "192.168.100.10", "443", dg="DG-ALL")
        self.assertNotEqual(getattr(r.rule, "name", None), "DGHN-PRE")

    def test_canh_bao_shadow_tu_rule_deny_o_shared_pre(self):
        r = _one("10.66.66.66", "192.168.100.10", "443")
        self.assertEqual(r.shadow_warning, "SHARED-PRE-DENY")
        self.assertEqual(_compute_action_level(r), "add")

    def test_token_khong_doc_duoc_van_bao_cao(self):
        r = _one("10.1.1.100-ABC", "192.168.100.10", "443")
        self.assertEqual(r.tier, "INVALID")
        self.assertIn("10.1.1.100-ABC", r.note)


class TestExportPanorama(unittest.TestCase):
    def setUp(self):
        self.out = tempfile.mkdtemp(prefix="checkrule_pan_")

    def tearDown(self):
        shutil.rmtree(self.out, ignore_errors=True)

    def test_file_excel_co_cot_vi_tri_rule_dung_vi_tri(self):
        cfg, idx, ris, zl = _ctx()
        pf = build_manual_request([
            ["10.7.7.7", "192.168.100.10", "443"],
            ["10.5.5.100-10.5.5.110", "192.168.100.10", "443"],
            ["10.20.20.20", "192.168.100.10", "443"],
            ["10.99.99.99", "192.168.100.10", "443"],
        ])
        res = [classify_row(r, ris, idx, zone_lookup=zl) for r in pf.rows]
        path = os.path.join(self.out, "kq.xlsx")
        export_results(pf.headers, res, path)

        ws = openpyxl.load_workbook(path)["Ket qua kiem tra"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[:5], ["Name Rule", "Vi tri Rule", "Source", "Dest", "Port"])
        vi_tri = [ws.cell(row=i, column=2).value for i in range(2, 6)]
        self.assertEqual(vi_tri[:3], ["DG-HN-PRE", "DG-DC-PRE", "DG-ALL-PRE"])
        # Dong TOP4/near_match cung phai co vi tri cua rule GOI Y (khong de trong)
        self.assertTrue(vi_tri[3])

    def test_sheet_chu_thich_noi_ro_2_tang_commit_cua_panorama(self):
        cfg, idx, ris, zl = _ctx()
        pf = build_manual_request([["10.7.7.7", "192.168.100.10", "443"]])
        res = [classify_row(r, ris, idx, zone_lookup=zl) for r in pf.rows]
        path = os.path.join(self.out, "kq2.xlsx")
        export_results(pf.headers, res, path)
        ws = openpyxl.load_workbook(path)["Chu thich"]
        text = " ".join(str(c.value or "") for row in ws.iter_rows() for c in row)
        self.assertIn("Vi tri Rule", text)
        self.assertIn("Commit and Push", text)
        self.assertIn("Shared-PRE", text)


if __name__ == "__main__":
    unittest.main()
