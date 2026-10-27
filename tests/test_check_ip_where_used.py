# -*- coding: utf-8 -*-
"""
Luoi chan cho muc 5 CHECK IP (ban PANORAMA) - "Check Object" (Where Used).

Tra loi: 1 IP (hoac 1 ten object) dang nam trong Address Object nao, Group
nao, va rule nao dung toi chung.

Bon cho PAN-OS khac han Check Point va deu de sai, moi cho mot nhom test:
  1. IP go THANG vao o Source cua rule (khong qua Address Object) - van phai
     tinh la THAM CHIEU.
  2. FQDN Object - "khong biet", khong phai "khong".
  3. Dynamic Address Group - thanh vien tinh LUC CHAY tu registered-ip tag.
  4. Chuoi Device Group (Shared -> DG cha -> DG dang chon) - sua nham TANG la
     khong co tac dung gi.

Rieng ban Panorama con kiem them: ket qua phai mang theo VI TRI cua rule
(Shared-PRE / <DG>-PRE / <DG>-POST / Shared-POST).

Khong can tkinter, khong can firewall.
"""

import unittest

from check_ip_pan.core import AddressGroup, AddressObject, NatRule, SecurityRule
from check_ip_pan.resolver_engine import GroupResolver
from check_ip_pan.where_used import (KIND_ANY, KIND_REF, KIND_VALUE,
                                 build_parent_index, format_report, walk_up,
                                 where_used_ip, where_used_object)

# --- kho object: co du ca 4 cho de sai ------------------------------------
OBJECTS = {
    "H-WEB": AddressObject("H-WEB", "ip-netmask", "10.1.1.5/32"),
    "H-WEB-NOMASK": AddressObject("H-WEB-NOMASK", "ip-netmask", "10.1.1.5"),
    "NET-10": AddressObject("NET-10", "ip-netmask", "10.0.0.0/8"),
    "NET-LAN": AddressObject("NET-LAN", "ip-netmask", "10.1.1.0/24"),
    "RANGE-DB": AddressObject("RANGE-DB", "ip-range", "10.1.1.1-10.1.1.9"),
    "FQ-PORTAL": AddressObject("FQ-PORTAL", "fqdn", "portal.vnpt.vn"),
    "H-OTHER": AddressObject("H-OTHER", "ip-netmask", "192.168.9.9/32"),
}
GROUPS = {
    "G-WEB": AddressGroup("G-WEB", "static", members=["H-WEB", "H-OTHER"]),
    "G-ALL": AddressGroup("G-ALL", "static", members=["G-WEB"]),
    "G-SUBNET": AddressGroup("G-SUBNET", "static", members=["NET-LAN"]),
    "DAG-PROD": AddressGroup("DAG-PROD", "dynamic", filter_expr="'prod'"),
}


def rule(index, name, source, dest, action="allow", disabled=False):
    return SecurityRule(index=index, name=name, source=list(source),
                        destination=list(dest), service=["any"],
                        action=action, disabled=disabled)


RULES = [
    rule(1, "REF-OBJECT", ["H-WEB"], ["any"]),
    rule(2, "REF-GROUP-LONG", ["G-ALL"], ["any"]),
    rule(3, "REF-LITERAL", ["10.1.1.5"], ["any"]),
    rule(4, "VALUE-SUBNET", ["NET-10"], ["any"]),
    rule(5, "VALUE-LITERAL-CIDR", ["10.1.1.0/24"], ["any"]),
    rule(6, "CHI-ANY", ["any"], ["any"]),
    rule(7, "KHONG-LIEN-QUAN", ["H-OTHER"], ["H-OTHER"]),
    rule(8, "REF-O-DESTINATION", ["H-OTHER"], ["H-WEB"]),
]

NAT_RULES = [
    NatRule(index=1, name="NAT-HIDE", source=["H-WEB"], destination=["any"],
            src_xlate_type="dynamic-ip-and-port", src_xlate_addresses=["H-OTHER"]),
    NatRule(index=2, name="NAT-DST", source=["any"], destination=["H-OTHER"],
            dst_xlate_type="static", dst_xlate_address="H-WEB"),
    NatRule(index=3, name="NAT-KHONG-LIEN-QUAN", source=["H-OTHER"],
            destination=["H-OTHER"]),
]


def _resolver(tags=None):
    return GroupResolver(OBJECTS, GROUPS, tags or {})


def _report(ip="10.1.1.5", tags=None, nat=None):
    return where_used_ip(ip, OBJECTS, GROUPS, RULES, _resolver(tags),
                         nat_rules=nat)


class TestTimObject(unittest.TestCase):
    def test_object_neu_dich_danh_va_object_chua_theo_subnet_duoc_TACH_RO(self):
        rep = _report()
        exact = {h.name for h in rep.objects if h.exact}
        wide = {h.name for h in rep.objects if not h.exact}
        self.assertIn("H-WEB", exact)
        self.assertIn("H-WEB-NOMASK", exact, "ip-netmask khong co '/' la host don")
        self.assertIn("RANGE-DB", exact, "ip-range luon la neu dich danh")
        self.assertEqual(wide, {"NET-10", "NET-LAN"})
        self.assertNotIn("H-OTHER", {h.name for h in rep.objects})

    def test_FQDN_la_KHONG_BIET_chu_khong_phai_KHONG(self):
        """FQDN khong resolve ra IP duoc. Im lang bo qua thi nguoi doc tuong
        tool da kiem tra xong va ket luan 'khong nam trong do'."""
        rep = _report()
        self.assertNotIn("FQ-PORTAL", {h.name for h in rep.objects})
        self.assertTrue(any("FQDN" in n and "FQ-PORTAL" in n for n in rep.notes),
                        rep.notes)

    def test_IP_khong_co_object_nao_mang_thi_NOI_RA(self):
        rep = _report(ip="172.16.9.9")
        self.assertEqual(rep.objects, [])
        self.assertTrue(rep.notes)

    def test_IP_hong_thi_noi_ro_chu_khong_tra_rong(self):
        rep = _report(ip="khong-phai-ip")
        self.assertTrue(rep.empty)
        self.assertTrue(any("khong phai dia chi" in n for n in rep.notes), rep.notes)


class TestCayGroup(unittest.TestCase):
    def test_di_nguoc_len_nhieu_cap(self):
        rep = _report()
        depth = {g.name: g.depth for g in rep.groups}
        self.assertEqual(depth.get("G-WEB"), 1)
        self.assertEqual(depth.get("G-ALL"), 2)

    def test_group_chua_SUBNET_rong_van_duoc_in_ra(self):
        """Nguoi van hanh can biet ca cac group chua subnet trum len IP nay -
        bo di la bo mat mot nua cau tra loi 'phai sua o dau'."""
        rep = _report()
        self.assertIn("G-SUBNET", {g.name for g in rep.groups})

    def test_long_vong_khong_treo(self):
        parents = {"A": {"B"}, "B": {"C"}, "C": {"A"}}
        got = walk_up(["A"], parents)
        self.assertEqual(sorted(g.name for g in got), ["B", "C"])

    def test_chi_muc_nguoc_chi_lay_group_STATIC(self):
        """Group dynamic KHONG co member trong config - dua vao chi muc nguoc
        la bia ra quan he khong ton tai."""
        parents = build_parent_index(GROUPS)
        self.assertIn("G-WEB", parents.get("H-WEB", set()))
        self.assertNotIn("DAG-PROD", {p for ps in parents.values() for p in ps})


class TestBaNhomRule(unittest.TestCase):
    def _names(self, rep):
        return ([h.rule_name for h in rep.rules_by_reference],
                [h.rule_name for h in rep.rules_by_value],
                [h.rule_name for h in rep.rules_by_any])

    def test_tach_rieng_THAM_CHIEU_GIA_TRI_va_ANY(self):
        ref, val, anyb = self._names(_report())
        self.assertIn("REF-OBJECT", ref)
        self.assertIn("REF-GROUP-LONG", ref, "rule goi group long phai la tham chieu")
        self.assertIn("REF-O-DESTINATION", ref)
        self.assertIn("VALUE-SUBNET", val)
        self.assertIn("VALUE-LITERAL-CIDR", val)
        self.assertIn("CHI-ANY", anyb)
        self.assertNotIn("KHONG-LIEN-QUAN", ref + val + anyb)

    def test_IP_GO_THANG_vao_rule_la_THAM_CHIEU(self):
        """PAN-OS cho go IP truc tiep vao o Source, khong can Address Object.
        Do la 'neu dich danh' - xep xuong nhom theo gia tri la sai han."""
        ref, val, _ = self._names(_report())
        self.assertIn("REF-LITERAL", ref)
        self.assertNotIn("REF-LITERAL", val)
        hit = next(h for h in _report().rules_by_reference
                   if h.rule_name == "REF-LITERAL")
        self.assertIn("go thang", hit.via)

    def test_subnet_rong_KHONG_duoc_tinh_la_tham_chieu(self):
        """Gop vao thi moi IP deu 'duoc rule 10.0.0.0/8 tham chieu' va ca 3
        nhom mat het y nghia."""
        ref, _, _ = self._names(_report())
        self.assertNotIn("VALUE-SUBNET", ref)

    def test_mot_rule_khong_bao_gio_nam_o_HAI_nhom(self):
        rep = _report()
        keys = [h.rule_index for h in rep.rules_by_reference] \
            + [h.rule_index for h in rep.rules_by_value] \
            + [h.rule_index for h in rep.rules_by_any]
        self.assertEqual(len(keys), len(set(keys)), "rule bi dem hai lan")

    def test_rule_vua_co_ten_vua_co_any_duoc_xep_theo_o_MANH_NHAT(self):
        """Rule 1 co Source=H-WEB, Destination=any. Xep no vao nhom 'chi vi any'
        la chon dung thong tin vo nghia nhat de dan nhan."""
        rep = _report()
        hit = next(h for h in rep.rules_by_reference if h.rule_name == "REF-OBJECT")
        self.assertEqual(hit.kind, KIND_REF)
        self.assertIn("Source", hit.sides)

    def test_bao_dung_ben_khop(self):
        rep = _report()
        hit = next(h for h in rep.rules_by_reference
                   if h.rule_name == "REF-O-DESTINATION")
        self.assertIn("Destination", hit.sides)
        self.assertNotIn("Source", hit.sides)

    def test_nhom_any_duoc_dem_va_noi_ra_trong_ghi_chu(self):
        rep = _report()
        self.assertTrue(any("'any'" in n for n in rep.notes), rep.notes)


class TestDynamicAddressGroup(unittest.TestCase):
    def test_co_tag_thi_tim_duoc_thanh_vien_va_noi_ro_la_LUC_CHAY(self):
        rep = _report(tags={"10.1.1.5": {"prod"}, "10.9.9.9": {"dev"}})
        dag = [g for g in rep.groups if g.group_kind == "dynamic"]
        self.assertEqual([g.name for g in dag], ["DAG-PROD"])
        self.assertIn("LUC CHAY", dag[0].note)

    def test_IP_khong_co_tag_khop_thi_khong_vao_DAG(self):
        rep = _report(tags={"10.1.1.5": {"dev"}})
        self.assertEqual([g for g in rep.groups if g.group_kind == "dynamic"], [])

    def test_KHONG_lay_duoc_tag_thi_phai_noi_la_KHONG_BIET(self):
        """Day la luc ket qua 'khong thay dung o dau' nguy hiem nhat: no trong
        y het su that."""
        rep = _report(tags=None)
        self.assertTrue(any("Dynamic Address Group" in n and "DAG-PROD" in n
                            for n in rep.notes), rep.notes)

    def test_rule_goi_DAG_dang_chua_IP_la_THAM_CHIEU(self):
        rules = RULES + [rule(9, "REF-DAG", ["DAG-PROD"], ["any"])]
        rep = where_used_ip("10.1.1.5", OBJECTS, GROUPS, rules,
                            _resolver({"10.1.1.5": {"prod"}}))
        self.assertIn("REF-DAG", [h.rule_name for h in rep.rules_by_reference])


class TestNat(unittest.TestCase):
    def test_quet_CA_O_DA_DICH(self):
        """O da dich moi tra loi 'IP nay ra ngoai thanh gi' - va do la o hay bi
        bo sot nhat khi di tim 'con nay dang NAT o dau'."""
        rep = _report(nat=NAT_RULES)
        names = {h.rule_name for h in rep.nat_hits}
        self.assertIn("NAT-HIDE", names)
        self.assertIn("NAT-DST", names, "thieu o Destination Translated")
        self.assertNotIn("NAT-KHONG-LIEN-QUAN", names)
        dst = next(h for h in rep.nat_hits if h.rule_name == "NAT-DST")
        self.assertIn("Destination Translated", dst.sides)

    def test_khong_truyen_NAT_thi_khong_no(self):
        self.assertEqual(_report(nat=None).nat_hits, [])


class TestTraTheoTenObject(unittest.TestCase):
    def test_dung_nghia_Where_Used_KHONG_suy_dien_theo_gia_tri(self):
        rep = where_used_object("H-WEB", OBJECTS, GROUPS, RULES, _resolver(),
                                nat_rules=NAT_RULES)
        ref = [h.rule_name for h in rep.rules_by_reference]
        self.assertIn("REF-OBJECT", ref)
        self.assertIn("REF-GROUP-LONG", ref)
        self.assertEqual(rep.rules_by_value, [],
                         "tra theo ten ma tra ve ca rule 'trum len gia tri' la "
                         "tra ve thu nguoi hoi khong hoi")
        self.assertNotIn("REF-LITERAL", ref,
                         "IP go thang khong phai tham chieu TOI OBJECT nay")

    def test_tra_duoc_ca_ten_GROUP(self):
        rep = where_used_object("G-WEB", OBJECTS, GROUPS, RULES, _resolver())
        self.assertIn("G-ALL", {g.name for g in rep.groups})
        self.assertIn("REF-GROUP-LONG", [h.rule_name for h in rep.rules_by_reference])

    def test_ten_khong_ton_tai_thi_NOI_RO(self):
        rep = where_used_object("KHONG-CO-THAT", OBJECTS, GROUPS, RULES, _resolver())
        self.assertTrue(rep.empty)
        self.assertTrue(any("Khong co Address Object/Group" in n for n in rep.notes))

    def test_object_KHONG_AI_DUNG_thi_ket_luan_duoc_la_XOA_DUOC(self):
        """He thong chi kiem soat theo IP/port nen Security + NAT la TOAN BO
        cho mot object co the duoc dung - "khong thay dung o dau" la cau tra
        loi DAY DU, phai noi thang ra thay vi im lang tra danh sach rong."""
        from check_ip_pan.where_used import UNUSED_TAIL

        objs = {"RAC": AddressObject("RAC", "ip-netmask", "10.99.99.99/32")}
        rep = where_used_object("RAC", objs, {}, RULES, GroupResolver(objs, {}, {}))
        self.assertEqual(rep.rules_by_reference, [])
        joined = " ".join(rep.notes)
        self.assertIn("RAC, xoa duoc", joined)
        self.assertIn(UNUSED_TAIL[:40], joined)

    def test_con_nam_trong_GROUP_thi_KHONG_duoc_bao_la_xoa_duoc(self):
        """PAN-OS tu choi xoa object dang con bi tham chieu - bao xoa duoc la
        day nguoi dung vao mot thao tac chac chan loi."""
        objs = {"RAC": AddressObject("RAC", "ip-netmask", "10.99.99.99/32")}
        grps = {"G-RAC": AddressGroup("G-RAC", "static", members=["RAC"])}
        rep = where_used_object("RAC", objs, grps, RULES,
                                GroupResolver(objs, grps, {}))
        joined = " ".join(rep.notes)
        self.assertNotIn("xoa duoc", joined)
        self.assertIn("go khoi group", joined)

    def test_canh_bao_chi_noi_DUNG_cho_tool_chua_do(self):
        """Canh bao phai la mot cau DO DUOC, khong phai cau ra dieu kien.

        Ban truoc con them "hoac o rule CUC BO tren chinh firewall (ngoai tam
        quan ly Panorama)". Nguoi van hanh bac bo: o he thong nay Panorama la
        tap CHA - cai gi firewall co thi Panorama deu co. Menh de do vua sai
        vua lam hong ket luan, vi bien mot cau tra loi do duoc thanh mot cau
        khong ai kiem chung duoc.
        """
        from check_ip_pan.where_used import UNUSED_WARN_DG

        self.assertIn("Device Group KHAC", UNUSED_WARN_DG)
        self.assertNotIn("CUC BO", UNUSED_WARN_DG,
                         "van con menh de 'rule cuc bo tren firewall' da bi bac bo")
        self.assertNotIn("firewall", UNUSED_WARN_DG.lower())

    def test_object_o_tang_SHARED_thi_canh_bao_MANH_HON(self):
        """Object o Shared/DG cha thi MOI nhanh DG con deu nhin thay."""
        objs = {"RAC": AddressObject("RAC", "ip-netmask", "10.99.99.99/32")}
        r = GroupResolver(objs, {}, {})
        xa = where_used_object("RAC", objs, {}, RULES, r,
                               layer_of={"RAC": ["shared", "DG-HN"]})
        self.assertIn("CAN THAN", " ".join(xa.notes))
        gan = where_used_object("RAC", objs, {}, RULES, r,
                                layer_of={"RAC": ["DG-HN-01"]})
        self.assertNotIn("CAN THAN", " ".join(gan.notes))
        self.assertIn("Device Group KHAC", " ".join(gan.notes))

    def test_object_DANG_duoc_dung_thi_khong_bao_xoa_duoc(self):
        rep = where_used_object("H-WEB", OBJECTS, GROUPS, RULES, _resolver())
        self.assertTrue(rep.rules_by_reference)
        self.assertNotIn("xoa duoc", " ".join(rep.notes))


class TestChuoiDeviceGroup(unittest.TestCase):
    """RIENG BAN PANORAMA: cung 1 ten co the ton tai o nhieu tang."""

    def test_object_bi_tang_GAN_HON_override_phai_duoc_canh_bao(self):
        """Sua ban o tang xa se KHONG doi gi o DG nay - khong noi ra thi nguoi
        dung sua dung ten, dung gia tri, nhung sai TANG."""
        rep = where_used_object("H-WEB", OBJECTS, GROUPS, RULES, _resolver(),
                                layer_of={"H-WEB": ["shared", "DG-HN", "DG-HN-01"]})
        note = " ".join(rep.notes)
        self.assertIn("override", note)
        self.assertIn("DG-HN-01", note, "phai noi ro tang nao dang THANG")
        self.assertIn("shared", note, "phai noi ro tang nao sua se vo ich")

    def test_object_chi_o_shared_thi_canh_bao_anh_huong_MOI_DG(self):
        rep = where_used_object("H-WEB", OBJECTS, GROUPS, RULES, _resolver(),
                                layer_of={"H-WEB": ["shared"]})
        self.assertTrue(any("MOI Device Group" in n for n in rep.notes), rep.notes)

    def test_object_chi_o_DG_hien_tai_thi_khong_canh_bao_thua(self):
        rep = where_used_object("H-WEB", OBJECTS, GROUPS, RULES, _resolver(),
                                layer_of={"H-WEB": ["DG-HN-01"]})
        self.assertFalse(any("override" in n or "MOI Device Group" in n
                             for n in rep.notes), rep.notes)


class TestViTriRule(unittest.TestCase):
    """RIENG BAN PANORAMA: moi rule thuoc 1 trong 4 vi tri, va thu tu do quyet
    dinh rule nao an rule nao."""

    def test_ket_qua_mang_theo_VI_TRI_cua_rule(self):
        rules = [rule(1, "R-SHARED-PRE", ["H-WEB"], ["any"])]
        rules[0].vi_tri = "Shared-PRE"
        rep = where_used_ip("10.1.1.5", OBJECTS, GROUPS, rules, _resolver())
        self.assertEqual(rep.rules_by_reference[0].vi_tri, "Shared-PRE")

    def test_IN_RA_man_hinh_co_vi_tri(self):
        """Thieu vi tri la thieu dung thu nguoi doc can de biet sua o dau."""
        rules = [rule(1, "R-POST", ["H-WEB"], ["any"])]
        rules[0].vi_tri = "DG-HN-POST"
        rep = where_used_ip("10.1.1.5", OBJECTS, GROUPS, rules, _resolver())
        self.assertIn("@DG-HN-POST", "\n".join(format_report(rep)))


class TestInRaManHinh(unittest.TestCase):
    def test_in_du_ca_ba_nhom_va_ghi_chu(self):
        """Yeu cau cua nguoi dung: ket qua phai doc duoc NGAY tren man hinh,
        file Excel chi la ban luu."""
        lines = format_report(_report(nat=NAT_RULES))
        blob = "\n".join(lines)
        self.assertIn("10.1.1.5", blob)
        self.assertIn("THEO THAM CHIEU", blob)
        self.assertIn("THEO GIA TRI", blob)
        self.assertIn("NAT Rule", blob)
        self.assertIn("(!)", blob, "ghi chu phai duoc in ra, khong chi nam trong file")

    def test_cat_bot_khi_qua_dai_nhung_NOI_RO_con_bao_nhieu(self):
        many = [rule(100 + i, f"R{i}", ["H-WEB"], ["any"]) for i in range(40)]
        rep = where_used_ip("10.1.1.5", OBJECTS, GROUPS, many, _resolver())
        blob = "\n".join(format_report(rep, max_rules=5))
        self.assertIn("va 35 rule nua", blob)


class TestKhongCatImLang(unittest.TestCase):
    """IP nam trong RAT NHIEU object/group - phai hien DU, giong Global Find.

    Day la ca hay gap nhat trong cau hinh that: 1 IP vua la host object rieng,
    vua nam trong 2-3 subnet object, vua nam trong hang chuc group long nhau.
    Cat bot ma khong noi ra thi nguoi doc tuong da thay het roi di ket luan -
    sai kieu nguy hiem nhat cua mot tinh nang tra cuu.
    """

    @staticmethod
    def _kho_lon(n_obj=60, n_grp=40):
        """Kho co n_obj object cung mang 10.1.1.5 va n_grp group long nhau."""
        objs = {f"H-{i:03d}": AddressObject(f"H-{i:03d}", "ip-netmask", "10.1.1.5/32")
                for i in range(n_obj)}
        grps = {"G-000": AddressGroup("G-000", "static",
                                      members=sorted(objs)[:5])}
        for i in range(1, n_grp):
            grps[f"G-{i:03d}"] = AddressGroup(f"G-{i:03d}", "static",
                                              members=[f"G-{i - 1:03d}"])
        return objs, grps

    def test_TIM_DU_khong_bo_sot_object_nao(self):
        objs, grps = self._kho_lon()
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}))
        self.assertEqual(len(rep.objects), 60, "bo sot object mang IP nay")

    def test_TIM_DU_group_long_nhau_trong_tran_do_sau(self):
        from check_ip_pan.where_used import MAX_GROUP_DEPTH

        objs, grps = self._kho_lon(n_grp=MAX_GROUP_DEPTH - 2)
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}))
        self.assertEqual(len([g for g in rep.groups if not g.note]),
                         MAX_GROUP_DEPTH - 2, "bo sot group cha")

    def test_cham_TRAN_do_sau_thi_PHAI_canh_bao(self):
        """Khong canh bao thi ket qua thieu ma trong y het ket qua du."""
        from check_ip_pan.where_used import MAX_GROUP_DEPTH

        objs, grps = self._kho_lon(n_grp=MAX_GROUP_DEPTH + 5)
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}))
        self.assertTrue(any("cham tran" in (g.note or "") for g in rep.groups),
                        "cat o tran 12 cap ma khong noi gi")

    def test_IN_RA_man_hinh_noi_ro_con_bao_nhieu_chua_hien(self):
        objs, grps = self._kho_lon(n_obj=300)
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}))
        blob = "\n".join(format_report(rep, max_items=200))
        self.assertIn("Address Object (300)", blob, "phai in TONG SO that")
        self.assertIn("va 100 object nua", blob, "cat bot ma khong noi ra")

    def test_group_cat_bot_cung_phai_noi_ra(self):
        objs, grps = self._kho_lon(n_obj=5, n_grp=11)
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}))
        blob = "\n".join(format_report(rep, max_items=4))
        self.assertIn("Address Group (", blob)
        self.assertIn("group nua", blob)

    def test_NAT_cat_bot_cung_phai_noi_ra(self):
        objs, grps = self._kho_lon(n_obj=1, n_grp=1)
        nats = [NatRule(index=i, name=f"N{i}", source=["H-000"], destination=["any"])
                for i in range(30)]
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}),
                            nat_rules=nats)
        blob = "\n".join(format_report(rep, max_rules=10))
        self.assertIn("va 20 NAT rule nua", blob)

    def test_via_nhieu_ten_thi_noi_ro_con_bao_nhieu(self):
        """O Source goi 10 group cung chua IP - khong duoc im lang chi in 3."""
        objs = {"H-A": AddressObject("H-A", "ip-netmask", "10.1.1.5/32")}
        grps = {f"G-{i}": AddressGroup(f"G-{i}", "static", members=["H-A"])
                for i in range(10)}
        r = rule(1, "R", sorted(grps), ["any"])
        rep = where_used_ip("10.1.1.5", objs, grps, [r], GroupResolver(objs, grps, {}))
        via = rep.rules_by_reference[0].via
        self.assertIn("+", via, f"cat im lang danh sach 'qua': {via}")
        self.assertIn("nua", via)

    def test_FILE_XUAT_thi_KHONG_cat_gi_het(self):
        """Man hinh cat cho gon duoc, file thi khong - do la ban day du."""
        import tempfile

        import openpyxl

        from check_ip_pan.excel_export import export_where_used_to_excel

        objs, grps = self._kho_lon(n_obj=300, n_grp=3)
        rep = where_used_ip("10.1.1.5", objs, grps, [], GroupResolver(objs, grps, {}))
        path = export_where_used_to_excel(tempfile.mkdtemp(), [rep])
        ws = openpyxl.load_workbook(path).worksheets[0]
        names = {ws.cell(row=r, column=3).value
                 for r in range(2, ws.max_row + 1)
                 if ws.cell(row=r, column=2).value == "OBJECT"}
        self.assertEqual(len(names), 300, "file xuat bi cat bot object")


class TestXuatFile(unittest.TestCase):
    def test_sheet_co_cot_VI_TRI_rieng_cua_Panorama(self):
        import tempfile

        import openpyxl

        from check_ip_pan.excel_export import (WHERE_HEADERS,
                                               export_where_used_to_excel)

        rules = [rule(1, "R-PRE", ["H-WEB"], ["any"])]
        rules[0].vi_tri = "Shared-PRE"
        rep = where_used_ip("10.1.1.5", OBJECTS, GROUPS, rules, _resolver())
        path = export_where_used_to_excel(tempfile.mkdtemp(), [rep])
        ws = openpyxl.load_workbook(path).worksheets[0]
        head = [ws.cell(row=1, column=c).value for c in range(1, len(WHERE_HEADERS) + 1)]
        self.assertIn("Vi tri Panorama", head)
        col = head.index("Vi tri Panorama") + 1
        found = {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
        self.assertIn("Shared-PRE", found)

    def test_sheet_giu_nguyen_BA_NHOM(self):
        import tempfile

        import openpyxl

        from check_ip_pan.excel_export import export_where_used_to_excel

        folder = tempfile.mkdtemp()
        path = export_where_used_to_excel(folder, [_report(nat=NAT_RULES)])
        ws = openpyxl.load_workbook(path).worksheets[0]
        groups = {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
        self.assertIn("RULE (tham chieu)", groups)
        self.assertIn("RULE (theo gia tri)", groups)
        self.assertIn("RULE (chi vi 'any')", groups)
        self.assertIn("NAT RULE", groups)
        self.assertIn("GHI CHU", groups)

    def test_csv_co_BOM_cho_Excel(self):
        import io as _io
        import tempfile

        from check_ip_pan.excel_export import export_where_used_to_csv

        path = export_where_used_to_csv(tempfile.mkdtemp(), [_report()])
        with _io.open(path, "rb") as f:
            self.assertEqual(f.read(3), b"\xef\xbb\xbf")


if __name__ == "__main__":
    unittest.main(verbosity=2)
