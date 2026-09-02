# -*- coding: utf-8 -*-
"""
Test excel_io.py (Mục 9 / Mục 6): 14 cot nen chung (D7), dinh dang Port D8
(UDP/`<port>`; TCP chi so port; any/application-default giu nguyen), export
xlsx/csv, red highlight khi KHAC ACTION, Import (thieu cot -> loi ro rang,
dong rong -> bo qua, doc lai dung danh sach (name, type)).
"""

import csv
import os
import shutil
import tempfile
import unittest
from datetime import datetime
import sys
try:
    import check_audit  # noqa
except ImportError:
    import check_audit_pan as check_audit  # type: ignore
    sys.modules['check_audit'] = check_audit
    for _sub in ['core','resolver','detectors','excel_io','pan_api']:
        try:
            _m = __import__(f'check_audit_pan.{_sub}', fromlist=[_sub])
            sys.modules[f'check_audit.{_sub}'] = _m
        except ImportError:
            pass

from check_audit.core import (
    Finding, ServiceObject, ServiceGroup,
    CHECK_TYPE_SHADOW, SHADOW_UNREACHABLE, SHADOW_UNUSED,
)
from check_audit import excel_io
from check_audit.excel_io import (
    HEADERS, export_findings, parse_import_file, _no_hit_duration_text,
)


def finding_shadow(name="r-shadow", service=None, same_action=None, reasons=None,
                   hit_count=None, last_hit_date=None, creation_date=None, no_hit_days=None):
    return Finding(
        check_type=CHECK_TYPE_SHADOW,
        rule_name=name, rule_uuid=f"uuid-{name}", rule_index=0,
        source=["10.0.0.1"], destination=["10.0.0.2"],
        service=list(service or ["TCP_1521"]), action="allow", disabled=False,
        detail="Bị che bởi rule 'R-master'", shadow_reasons=list(reasons or [SHADOW_UNREACHABLE]),
        same_action=same_action,
        hit_count=hit_count, last_hit_date=last_hit_date,
        rule_creation_date=creation_date, no_hit_days=no_hit_days,
    )


def finding_disabled(name="r-disabled", service=None):
    return Finding(
        check_type="disabled", rule_name=name, rule_uuid="", rule_index=2,
        source=["any"], destination=["any"], service=list(service or ["any"]),
        action="deny", disabled=True, detail="Rule đang disabled",
    )


class TestHeadersAndPortFormat(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        svc_objs = {
            "TCP_1521": ServiceObject("TCP_1521", "tcp", "1521"),
            "UDP_689": ServiceObject("UDP_689", "udp", "689"),
        }
        self.svc_info = excel_io.build_service_info_for_export(svc_objs, {})

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_headers_exact(self):
        f = finding_shadow(service=["TCP_1521"])
        path = export_findings([f], self.tmp, service_info=self.svc_info)
        last = os.listdir(self.tmp)[0]
        # kiem tra bang doc lai file that
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(self.tmp, last))
        ws = wb.active
        headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
        self.assertEqual(headers, HEADERS)
        # Pan has 15 cols (thêm Vi tri Rule), standalone 14
        self.assertIn(len(HEADERS), (14, 15))

    def test_port_d8_format(self):
        # D8: UDP -> "UDP/689"; TCP -> chi "1521"; any/application-default giu nguyen
        findings = [
            finding_shadow(name="a", service=["TCP_1521"]),
            finding_shadow(name="b", service=["UDP_689"]),
            finding_disabled(name="c", service=["any"]),
            finding_disabled(name="d", service=["application-default"]),
            finding_shadow(name="e", service=["TCP_1521", "UDP_689"]),
        ]
        path = export_findings(findings, self.tmp, types=["shadow", "disabled"], service_info=self.svc_info)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        # Pan has Vi tri at col3, so Port shifts from 5 to 6
        port_col = 6 if len(HEADERS)==15 else 5
        name_col = 2
        rows = {ws.cell(row=r, column=name_col).value: ws.cell(row=r, column=port_col).value
                for r in range(2, ws.max_row + 1)}
        self.assertEqual(rows["a"], "1521")
        self.assertEqual(rows["b"], "UDP/689")
        self.assertEqual(rows["c"], "any")
        self.assertEqual(rows["d"], "application-default")
        self.assertEqual(rows["e"], "1521, UDP/689")

    def test_csv_export_utf8(self):
        f = finding_shadow()
        path = export_findings([f], self.tmp, fmt="csv")
        self.assertTrue(path.endswith(".csv"))
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        self.assertEqual(rows[0], HEADERS)
        self.assertEqual(rows[1][1], "r-shadow")

    def test_red_font_action_when_khac_action(self):
        same = finding_shadow(name="same", same_action=True, reasons=[SHADOW_UNREACHABLE])
        diff = finding_shadow(name="diff", same_action=False, reasons=[SHADOW_UNREACHABLE])
        findings = [same, diff]
        export_findings(findings, self.tmp, service_info=self.svc_info)
        from openpyxl import load_workbook
        fname = os.listdir(self.tmp)[0]
        wb = load_workbook(os.path.join(self.tmp, fname))
        ws = wb.active
        # Action col shifts in Pan (Vi tri inserted)
        action_col = 7 if len(HEADERS)==15 else 6
        colors = {}
        for r in range(2, ws.max_row + 1):
            fc = ws.cell(row=r, column=action_col).font.color
            colors[ws.cell(row=r, column=2).value] = fc.rgb if fc else None
        self.assertEqual(colors["diff"], "FFC00000")
        self.assertNotEqual(colors["same"], "FFC00000")

    def test_no_hit_duration_text(self):
        f = finding_shadow(
            name="u1", same_action=True,
            reasons=[SHADOW_UNREACHABLE, SHADOW_UNUSED],
            hit_count=0, creation_date="2026-01-01", no_hit_days=45,
        )
        txt = _no_hit_duration_text(f)
        self.assertIn("45 ngày", txt)
        self.assertIn("Chưa từng có hit", txt)
        f2 = finding_shadow(name="u2", same_action=True, reasons=[SHADOW_UNREACHABLE])
        self.assertEqual(_no_hit_duration_text(f2), "")


class TestEmptyExport(unittest.TestCase):
    def test_raise_on_empty(self):
        with self.assertRaises(ValueError):
            export_findings([], "C:/whatever", types=["shadow"])


class TestImport(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_xlsx(self, headers, rows):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        path = os.path.join(self.tmp, "import_test.xlsx")
        wb.save(path)
        return path

    def test_missing_name_rule_raises(self):
        # Thieu cot bat buoc 'Name Rule' -> loi ro rang, khong doan
        path = self._make_xlsx(
            ["Loại (Finding Type)", "Source IP"],
            [["shadow", "10.0.0.1"], ["shadow", "10.0.0.2"]],
        )
        with self.assertRaises(ValueError) as ctx:
            parse_import_file(path)
        self.assertIn("Name Rule", str(ctx.exception))

    def test_missing_loai_raises(self):
        path = self._make_xlsx(
            ["Name Rule", "Source IP"],
            [["r1", "10.0.0.1"]],
        )
        with self.assertRaises(ValueError) as ctx:
            parse_import_file(path)
        self.assertIn("Loại", str(ctx.exception))

    def test_full_roundtrip(self):
        # Export roi Import lai: dung danh sach (name, loai)
        f = finding_shadow(service=["TCP_1521"])
        svc_objs = {"TCP_1521": ServiceObject("TCP_1521", "tcp", "1521")}
        svc_info = excel_io.build_service_info_for_export(svc_objs, {})
        out_dir = os.path.join(self.tmp, "out")
        path = export_findings([f], out_dir, types=["shadow"], service_info=svc_info)
        rows = parse_import_file(path)
        # Pan returns 4-tuple (name, loai, vi_tri, hit), standalone 2-tuple
        if len(HEADERS)==15:
            self.assertEqual(rows[0][:2], ("r-shadow", "Shadow Rules"))
            self.assertEqual(len(rows[0]), 4)
        else:
            self.assertEqual(rows, [("r-shadow", "Shadow Rules")])

    def test_blank_name_row_skipped(self):
        # Pan requires Vi tri Rule
        headers = ["Name Rule", "Loại (Finding Type)", "Vi tri Rule"] if len(HEADERS)==15 else ["Name Rule", "Loại (Finding Type)"]
        rows_data = [["r1", "Shadow Rules", "DG1-PRE"], ["", "Disabled Rules", "DG1-PRE"], ["r2", "Disabled Rules", "DG1-PRE"]] if len(HEADERS)==15 else [["r1", "Shadow Rules"], ["", "Disabled Rules"], ["r2", "Disabled Rules"]]
        path = self._make_xlsx(headers, rows_data)
        logs = []
        rows = parse_import_file(path, log_func=logs.append)
        # Pan returns 4-tuple, standalone 2-tuple
        if len(HEADERS)==15:
            self.assertEqual([r[:2] for r in rows], [("r1", "Shadow Rules"), ("r2", "Disabled Rules")])
        else:
            self.assertEqual(rows, [("r1", "Shadow Rules"), ("r2", "Disabled Rules")])
        self.assertTrue(any("bo qua" in m.lower() for m in logs))


if __name__ == "__main__":
    unittest.main()