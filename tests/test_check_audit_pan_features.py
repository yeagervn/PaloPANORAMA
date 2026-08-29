# -*- coding: utf-8 -*-
"""
Test tính năng mới check_audit_pan (Panorama) theo yêu cầu:
- Shadow hit_count==0 -> Import chỉ Disable
- Disabled -> Import Delete
- Overly Permissive -> chỉ Export (Import bỏ qua)
- pan_api disable/delete xpath
- excel_io 15 cột + parse 4-tuple + Hit Count
- GUI logic phân biệt (mô phỏng)
"""
import os
import shutil
import tempfile
import unittest
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta

# Import pan modules
import sys
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from check_audit_pan.core import Finding, RuleSnapshot, RuleHitStat, ServiceObject, ServiceGroup
from check_audit_pan import excel_io as pan_excel
from check_audit_pan.excel_io import HEADERS, export_findings, parse_import_file
from check_audit_pan.pan_api import PanApiClient
from check_audit_pan.detectors import find_shadow_rules, find_disabled_rules, find_overly_permissive_rules
from check_audit_pan.resolver import GroupResolver, ServiceResolver

# helpers local
def mk_rule(index, name, action="allow", disabled=False, src=None, dst=None, svc=None, from_zones=None, to_zones=None):
    r = RuleSnapshot(
        index=index, name=name, uuid=f"uuid-{name}",
        from_zones=list(from_zones or ["any"]), to_zones=list(to_zones or ["any"]),
        source=list(src or ["any"]), destination=list(dst or ["any"]),
        service=list(svc or ["any"]), application=[],
        action=action, disabled=disabled, vi_tri="DG1-PRE", rulebase="pre"
    )
    # set resolved any for simplicity unless test overrides
    from check_audit_pan.resolver import ResolvedAddr, ResolvedService
    # any -> resolved any
    if r.source == ["any"] or r.source == []:
        r.resolved_source = ResolvedAddr(is_any=True)
    if r.destination == ["any"] or r.destination == []:
        r.resolved_destination = ResolvedAddr(is_any=True)
    if r.service == ["any"] or r.service == []:
        r.resolved_service = ResolvedService(is_any=True)
    return r

def hit_stat(hit_count=None, last_hit=None, creation=None):
    return RuleHitStat(hit_count, last_hit, creation)


class TestPanHeaders(unittest.TestCase):
    def test_headers_15_and_vi_tri(self):
        self.assertEqual(len(HEADERS), 15, "Panorama phải có 15 cột (thêm Vi tri Rule)")
        self.assertIn("Vi tri Rule", HEADERS)
        self.assertIn("Hit Count", HEADERS)
        self.assertEqual(HEADERS[1], "Name Rule")
        self.assertEqual(HEADERS[2], "Vi tri Rule")

    def test_export_has_vi_tri(self):
        tmp = tempfile.mkdtemp()
        try:
            f = Finding(check_type="shadow", rule_name="r1", rule_uuid="u1", rule_index=0,
                        source=["10.0.0.1"], destination=["10.0.0.2"], service=["TCP_1521"],
                        action="allow", disabled=False, detail="shadow",
                        shadow_reasons=["UNREACHABLE"], vi_tri="DG1-PRE")
            path = export_findings([f], tmp, types=["shadow"])
            from openpyxl import load_workbook
            wb = load_workbook(path)
            ws = wb.active
            headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column+1)]
            self.assertEqual(headers, HEADERS)
            # check vi_tri column value
            val = ws.cell(row=2, column=3).value
            self.assertEqual(val, "DG1-PRE")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestPanParseImport(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)
    def _make(self, headers, rows):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        p = os.path.join(self.tmp, "x.xlsx")
        wb.save(p)
        return p

    def test_parse_returns_4tuple_with_hit(self):
        path = self._make(HEADERS, [["Shadow Rules","r1","DG1-PRE","10.0.0.1","10.0.0.2","1521","allow","No","d","u","UNUSED","0","2026-01-01","2025-01-01","45 ngay"]])
        rows = parse_import_file(path)
        self.assertEqual(len(rows[0]), 4)
        self.assertEqual(rows[0][0], "r1")
        self.assertEqual(rows[0][1], "Shadow Rules")
        self.assertEqual(rows[0][2], "DG1-PRE")
        self.assertEqual(rows[0][3], "0")

    def test_missing_hit_column_compat(self):
        # file cũ không có Hit Count -> vẫn parse được, hit_raw = ""
        headers_old = ["Loại (Finding Type)","Name Rule","Vi tri Rule"]
        path = self._make(headers_old, [["Shadow Rules","r1","DG1-PRE"]])
        rows = parse_import_file(path)
        self.assertEqual(rows[0][3], "")

    def test_missing_vi_tri_raises(self):
        path = self._make(["Name Rule","Loại (Finding Type)"],[["r1","Shadow Rules"]])
        with self.assertRaises(ValueError) as ctx:
            parse_import_file(path)
        self.assertIn("Vi tri Rule", str(ctx.exception))

    def test_blank_name_skipped(self):
        path = self._make(HEADERS, [["Shadow Rules","","DG1-PRE","","","","","","","","","","","",""] , ["Shadow Rules","r2","DG1-PRE","","","","","","","","","","","",""]])
        logs=[]
        rows = parse_import_file(path, log_func=logs.append)
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0][0],"r2")

    def test_overly_parse_still_returns_but_gui_skips(self):
        path = self._make(HEADERS, [["Overly Permissive Rules","r-over","DG1-PRE","any","any","any","allow","No","d","","","","","",""]])
        rows = parse_import_file(path)
        self.assertEqual(rows[0][1], "Overly Permissive Rules")
        # GUI should skip overly - test logic below


class TestPanApiDisableDelete(unittest.TestCase):
    class FakeClient(PanApiClient):
        def __init__(self):
            super().__init__("1.1.1.1","key")
            self.calls=[]
        def _request(self, params):
            self.calls.append(dict(params))
            return ET.fromstring(b'<response status="success"/>')

    def test_delete_xpath_pre(self):
        c = self.FakeClient()
        c.delete_rule_candidate("DG1","pre","r1")
        p=c.calls[-1]
        self.assertEqual(p["action"],"delete")
        self.assertIn("pre-rulebase", p["xpath"])
        self.assertIn("DG1", p["xpath"])
        self.assertNotIn("disabled", p["xpath"])

    def test_disable_xpath_pre(self):
        c = self.FakeClient()
        c.disable_rule_candidate("DG1","pre","r1")
        p=c.calls[-1]
        self.assertEqual(p["action"],"set")
        self.assertIn("pre-rulebase", p["xpath"])
        self.assertIn("/disabled", p["xpath"])
        self.assertEqual(p["element"], "<disabled>yes</disabled>")

    def test_disable_shared(self):
        c = self.FakeClient()
        c.disable_rule_candidate("", "shared-pre","r1")
        p=c.calls[-1]
        self.assertIn("/config/shared/pre-rulebase", p["xpath"])
        self.assertEqual(p["action"],"set")

    def test_disable_post(self):
        c = self.FakeClient()
        c.disable_rule_candidate("DG2","post","r2")
        self.assertIn("post-rulebase", c.calls[-1]["xpath"])

    def test_delete_shared_post(self):
        c = self.FakeClient()
        c.delete_rule_candidate("", "shared-post","r3")
        self.assertIn("shared/post-rulebase", c.calls[-1]["xpath"])

    def test_no_commit_in_any(self):
        c = self.FakeClient()
        c.delete_rule_candidate("DG1","pre","r1")
        c.disable_rule_candidate("DG1","pre","r2")
        for call in c.calls:
            self.assertNotIn("commit", call.get("type",""))
            self.assertNotIn("commit", call.get("action",""))


class TestGuiLogicDifferentiation(unittest.TestCase):
    """Mô phỏng logic _on_import phân biệt action"""
    def _classify(self, rows, findings, rule_by_key):
        # mimic gui _on_import classification
        candidates=[]
        skipped_overly=0
        for row in rows:
            if len(row)==4:
                name, loai, vi_tri, hit_raw = row
            else:
                name, loai, vi_tri = row[:3]
                hit_raw=""
            if loai=="Overly Permissive Rules":
                skipped_overly+=1
                continue
            if loai not in ("Shadow Rules","Disabled Rules"):
                continue
            if loai=="Disabled Rules":
                action="delete"
            else:
                # shadow -> disable
                finding = next((f for f in findings if f.rule_name==name and f.vi_tri==vi_tri), None)
                # hit logic
                action="disable"
            dg, rb = ("DG1","pre")  # simplified
            candidates.append((name, loai, vi_tri, dg, rb, action))
        return candidates, skipped_overly

    def test_shadow_hit0_disable(self):
        findings=[
            Finding(check_type="shadow", rule_name="r-shadow", rule_uuid="u", rule_index=0, source=["10.0.0.1"], destination=["10.0.0.2"], service=["TCP_1521"], action="allow", disabled=False, detail="d", shadow_reasons=["UNUSED"], hit_count=0, vi_tri="DG1-PRE")
        ]
        rows=[("r-shadow","Shadow Rules","DG1-PRE","0")]
        cands, skip = self._classify(rows, findings, {})
        self.assertEqual(cands[0][5], "disable")
        self.assertEqual(skip,0)

    def test_disabled_delete(self):
        findings=[
            Finding(check_type="disabled", rule_name="r-dis", rule_uuid="u", rule_index=0, source=["any"], destination=["any"], service=["any"], action="deny", disabled=True, detail="d", vi_tri="DG1-PRE")
        ]
        rows=[("r-dis","Disabled Rules","DG1-PRE","")]
        cands,_ = self._classify(rows, findings, {})
        self.assertEqual(cands[0][5], "delete")

    def test_overly_skipped(self):
        rows=[("r-over","Overly Permissive Rules","DG1-PRE",""), ("r-shadow","Shadow Rules","DG1-PRE","0")]
        cands, skip = self._classify(rows, [], {})
        self.assertEqual(len(cands),1)
        self.assertEqual(skip,1)
        self.assertEqual(cands[0][0],"r-shadow")

    def test_mixed_batch(self):
        rows=[
            ("r-s1","Shadow Rules","DG1-PRE","0"),
            ("r-s2","Shadow Rules","DG1-PRE","5"),
            ("r-d1","Disabled Rules","DG1-PRE",""),
            ("r-o1","Overly Permissive Rules","DG1-PRE",""),
        ]
        cands, skip = self._classify(rows, [], {})
        self.assertEqual(len(cands),3)
        self.assertEqual(skip,1)
        actions=dict((name,act) for name,loai,vi,dg,rb,act in cands)
        self.assertEqual(actions["r-s1"],"disable")
        self.assertEqual(actions["r-s2"],"disable")
        self.assertEqual(actions["r-d1"],"delete")


class TestDetectorsNoRegression(unittest.TestCase):
    def test_shadow_still_detects(self):
        from check_audit_pan.resolver import ResolvedAddr
        r1 = mk_rule(0,"R1", src=["any"], dst=["any"], svc=["any"])
        r2 = mk_rule(1,"R2", src=["10.0.0.5"], dst=["10.0.0.6"], svc=["TCP_1521"])
        # need resolved for r2 host
        import ipaddress
        def ip_int(ip): return int(ipaddress.ip_address(ip))
        from check_audit_pan.resolver import ResolvedAddr, ResolvedService
        r2.resolved_source = ResolvedAddr(ranges=[(ip_int("10.0.0.5"), ip_int("10.0.0.5"))])
        r2.resolved_destination = ResolvedAddr(ranges=[(ip_int("10.0.0.6"), ip_int("10.0.0.6"))])
        r2.resolved_service = ResolvedService(ports=[("tcp",1521,1521)])
        out = find_shadow_rules([r1,r2], None, None, {}, datetime.now())
        self.assertEqual(len(out),1)
        self.assertEqual(out[0].rule_name,"R2")

    def test_disabled_still(self):
        r = mk_rule(0,"D", disabled=True)
        out = find_disabled_rules([r])
        self.assertEqual(len(out),1)

    def test_overly_still(self):
        r = mk_rule(0,"O", src=["any"], dst=["10.0.0.1"], svc=["TCP_1521"])
        out = find_overly_permissive_rules([r])
        self.assertEqual(len(out),1)


if __name__ == "__main__":
    unittest.main()
